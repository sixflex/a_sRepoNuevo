"""Importar históricos desde la plataforma (RF-MIG-01, RF-MIG-02)."""
from io import BytesIO, StringIO

import openpyxl
from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse
from unittest.mock import patch
from django.core.exceptions import ValidationError

from academico.models import Asignatura, Campus, Carrera, Facultad, Sede
from auditoria.models import AuditoriaCambio
from encuestas import services as encuestas_services
from encuestas.models import RespuestaFormulario
from socios.models import ClasificacionSocio, Comuna, ParticipacionSocio, SocioComunitario
from academico.models import Seccion
from proyectos.models import Equipo

from . import services
from .models import CargaHistorica

ALMACEN_EN_MEMORIA = {
    **settings.STORAGES,
    "private": {"BACKEND": "django.core.files.storage.InMemoryStorage"},
}

ENCABEZADOS_REGISTRO = [
    "Marca temporal",
    "1.Nombre completo de quien esta respondiendo esta encuesta",
    "2.Rut (sin puntos y con guion",
    "5.Sede",
    "7.Carrera",
    "9.Nombre del Socio Comunitario (Institución, Empresa o Comunidad)",
    "10.Representante Socio Comunitario",
    "11.Cargo del Representante Socio Comunitario",
    "12.Teléfono de contacto con el Socio Comunitario",
    "13.Correo electrónico del Socio Comunitarios",
    "14.Comuna a la que pertenece el Socio Comunitario",
    "15.Dirección",
    "16.Clasifique al Socio Comunitarios",
    "17.Año en el que trabajo con el SC",
    "18.Número de  BENEFICIARIOS (personas)  que se llegó a favorecer con el proyecto A+S (aproximadamente)",
    "20.Semestre en el que trabajó con el SC",
    "21.Evalúe al Socio Comunitarios (En cuanto a su recomendación para volver a trabajar con él)",
    "Columna extra",
]


def excel(encabezados, filas, nombre="respuestas.xlsx"):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Respuestas"
    hoja.append(encabezados)
    for fila in filas:
        hoja.append(fila)
    contenido = BytesIO()
    libro.save(contenido)
    return SimpleUploadedFile(
        nombre, contenido.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def solicitud(usuario):
    request = RequestFactory().post("/")
    request.user = usuario
    return request


@override_settings(STORAGES=ALMACEN_EN_MEMORIA)
class ImportarRegistroSociosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.coordinador = get_user_model().objects.create_user(username="coord_hist", password="x")
        cls.coordinador.groups.add(Group.objects.create(name="Coordinador"))
        sede = Sede.objects.create(nombre="Sede Santiago", activo=True)
        cls.campus = Campus.objects.create(nombre="Campus Providencia", sede=sede, activo=True)
        facultad = Facultad.objects.create(nombre="Ciencias Sociales y Humanidades", activo=True)
        Carrera.objects.create(nombre="Psicología", facultad=facultad, activo=True)
        Comuna.objects.create(nombre="Ñuñoa", region="RM", lat_centro=-33.45, lon_centro=-70.6, activo=True)
        ClasificacionSocio.objects.create(nombre="Fundación", activo=True)

    def filas(self):
        return [
            ["2024-05-02 10:00:00", "Ana Pérez", "11111111-1", "Santiago", "Psicología", "Fundación Las Rosas",
             "María Soto", "Directora", "+56911112222", "contacto@lasrosas.cl", "Nunoa", "Av. Grecia 123",
             "Fundación", "2023;2024", "40", "Primavera", "5", "x"],
            ["2024-05-03 11:00:00", "Luis Rojas", "", "Santiago", "Carrera inventada", "Junta Vecinal Norte",
             "", "", "", "", "Comuna inexistente", "", "Club raro", "2022", "muchos", "Otoño", "9", ""],
            ["2024-05-04 12:00:00", "Sin socio", "", "", "", "", "", "", "", "", "", "", "", "2022", "", "", "", ""],
        ]

    def test_vista_previa_no_importa_y_clasifica_filas(self):
        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()), tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )
        self.assertEqual(carga.estado, services.ESTADO_PREVIA)
        self.assertEqual((carga.total_filas, carga.aceptadas, carga.observadas, carga.omitidas), (3, 1, 1, 1))
        self.assertFalse(SocioComunitario.objects.exists())
        self.assertEqual(carga.hoja, "Respuestas")
        self.assertIn("Comuna", carga.mapeos.values_list("campo_destino", flat=True))
        observada = carga.filas.get(estado=services.FILA_OBSERVADA)
        codigos = set(observada.errores.values_list("codigo", flat=True))
        self.assertTrue({"COMUNA_NO_RECONOCIDA", "CARRERA_NO_RECONOCIDA", "FUERA_DE_ESCALA", "NUMERO_INVALIDO"} <= codigos)
        self.assertEqual(carga.filas.get(estado=services.FILA_OMITIDA).errores.get().codigo, "SIN_SOCIO")
        # Se conserva la fila tal como venía (RF-MIG-01 CA2).
        self.assertEqual(carga.filas.get(numero_fila=2).datos_raw_json["Columna extra"], "x")

    def test_confirmar_crea_socios_y_participaciones_historicas(self):
        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()), tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )
        creados = services.confirmar_carga(solicitud(self.coordinador), carga, self.campus)

        self.assertEqual(creados, {"socios_nuevos": 2, "participaciones": 3})
        socio = SocioComunitario.objects.get(nombre_organizacion="Fundación Las Rosas")
        self.assertEqual(socio.comuna.nombre, "Ñuñoa")
        self.assertEqual(socio.estado_revision, "TRABAJO_TERMINADO")
        self.assertEqual(socio.direccion_exacta, "Av. Grecia 123")
        self.assertEqual(socio.contactos.get().correo, "contacto@lasrosas.cl")
        participaciones = ParticipacionSocio.objects.filter(socio=socio).order_by("fecha_inicio")
        self.assertEqual([p.fecha_inicio.year for p in participaciones], [2023, 2024])
        self.assertTrue(all(p.origen == "HISTORICO" and p.seccion_id is None and p.equipo_id is None for p in participaciones))
        self.assertEqual(participaciones[0].fecha_inicio.month, 8)
        self.assertEqual(participaciones[0].campus, self.campus)
        otro = SocioComunitario.objects.get(nombre_organizacion="Junta Vecinal Norte")
        self.assertIsNone(otro.comuna)
        self.assertEqual(otro.clasificacion_otra, "Club raro")
        self.assertIsNone(otro.participaciones.get().evaluacion)

        carga.refresh_from_db()
        self.assertEqual(carga.estado, services.ESTADO_CONFIRMADA)
        self.assertEqual(carga.filas.filter(socio_resultante__isnull=False).count(), 2)
        self.assertTrue(AuditoriaCambio.objects.filter(entidad="CargaHistorica", entidad_id=str(carga.pk)).exists())

    def test_subir_el_mismo_archivo_otra_vez_omite_lo_importado(self):
        primera = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()), tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )
        services.confirmar_carga(solicitud(self.coordinador), primera, self.campus)
        segunda = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()), tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )
        self.assertEqual(segunda.omitidas, 3)
        self.assertEqual(set(segunda.filas.values_list("errores__codigo", flat=True)), {"DUPLICADA", "SIN_SOCIO"})

    
    def test_importacion_historica_no_crea_secciones_ni_equipos(self):
        """CDE-93: importar históricos no crea secciones ni equipos activos."""

        secciones_antes = Seccion.objects.count()
        equipos_antes = Equipo.objects.count()

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        self.assertEqual(Seccion.objects.count(), secciones_antes)
        self.assertEqual(Equipo.objects.count(), equipos_antes)

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )

        self.assertEqual(Seccion.objects.count(), secciones_antes)
        self.assertEqual(Equipo.objects.count(), equipos_antes)

        participaciones = ParticipacionSocio.objects.filter(
            origen="HISTORICO"
        )

        self.assertEqual(participaciones.count(), 3)

        for participacion in participaciones:
            self.assertIsNone(participacion.seccion_id)
            self.assertIsNone(participacion.equipo_id)

    def test_consulta_historicos_con_filtros_combinados(self):
        """CDE-94: comprobar filtros combinados de históricos."""

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )

        self.client.force_login(self.coordinador)
        url = reverse("historicos:consulta_historicos")

        respuesta = self.client.get(url)

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            respuesta.context["total_participaciones"], 3
        )

        respuesta = self.client.get(url, {"anio": "2024"})

        self.assertEqual(
            respuesta.context["total_participaciones"], 1
        )
        respuesta = self.client.get(url, {
            "anio": "2024",
            "sede": str(self.campus.sede_id),
            "socio": "Las Rosas",
        })

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            respuesta.context["total_participaciones"], 1
        )
        self.assertEqual(
            respuesta.context["total_socios"], 1
        )

        self.assertContains(respuesta, "Fundación Las Rosas")
        self.assertNotContains(respuesta, "Junta Vecinal Norte")
        respuesta = self.client.get(url, {
            "anio": "2022",
            "socio": "Las Rosas",
        })

        self.assertEqual(
            respuesta.context["total_participaciones"], 0
        )

    def test_consulta_historicos_filtra_por_carrera(self):
        """CDE-94: comprobar el filtro por carrera."""

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )

        carrera = Carrera.objects.get(nombre="Psicología")

        participaciones = ParticipacionSocio.objects.filter(
            origen="HISTORICO"
        ).order_by("id")

        participaciones.update(carrera=None)

        participacion = participaciones.first()
        self.assertIsNotNone(participacion)

        participacion.carrera = carrera
        participacion.save(update_fields=["carrera"])

        self.client.force_login(self.coordinador)

        respuesta = self.client.get(
            reverse("historicos:consulta_historicos"),
            {"carrera": str(carrera.id)},
        )

        self.assertEqual(respuesta.status_code, 200)

  
        self.assertEqual(
            respuesta.context["total_participaciones"], 1
        )

        self.assertEqual(
            respuesta.context["total_socios"], 1
        )

        resultados = list(respuesta.context["pagina"])

        self.assertEqual(len(resultados), 1)
        self.assertEqual(resultados[0].id, participacion.id)
        self.assertEqual(resultados[0].carrera_id, carrera.id)


    def test_consulta_historicos_filtra_por_asignatura(self):
        """CDE-94: comprobar el filtro por asignatura."""

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )

        asignatura = Asignatura.objects.create(
            codigo="TEST-AS-001",
            nombre="Asignatura Histórica de Prueba",
            activo=True,
        )
        participaciones = ParticipacionSocio.objects.filter(
            origen="HISTORICO"
        ).order_by("id")

        participaciones.update(asignatura=None)

        participacion = participaciones.first()
        self.assertIsNotNone(participacion)

        participacion.asignatura = asignatura
        participacion.save(update_fields=["asignatura"])

        self.client.force_login(self.coordinador)

        respuesta = self.client.get(
            reverse("historicos:consulta_historicos"),
            {"asignatura": str(asignatura.id)},
        )

        self.assertEqual(respuesta.status_code, 200)

        self.assertEqual(
            respuesta.context["total_participaciones"], 1
        )

        self.assertEqual(
            respuesta.context["total_socios"], 1
        )

        resultados = list(respuesta.context["pagina"])

        self.assertEqual(len(resultados), 1)
        self.assertEqual(resultados[0].id, participacion.id)
        self.assertEqual(
            resultados[0].asignatura_id,
            asignatura.id,
        )

        self.assertContains(
            respuesta,
            "Asignatura Histórica de Prueba",
        )

    def test_consulta_historicos_paginacion(self):
        """CDE-94: comprobar paginación y conservación de filtros."""

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )
        original = ParticipacionSocio.objects.filter(
            origen="HISTORICO"
        ).first()

        self.assertIsNotNone(original)

        for i in range(22):
            participacion = ParticipacionSocio.objects.get(
                pk=original.pk
            )
            participacion.pk = None
            participacion.save()

        self.client.force_login(self.coordinador)
        url = reverse("historicos:consulta_historicos")
        respuesta = self.client.get(url)

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            respuesta.context["total_participaciones"], 25
        )
        self.assertEqual(
            len(respuesta.context["pagina"]), 20
        )
        self.assertTrue(
            respuesta.context["pagina"].has_next()
        )
        respuesta = self.client.get(url, {"pagina": "2"})

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            len(respuesta.context["pagina"]), 5
        )
        respuesta = self.client.get(url, {
            "socio": "Las Rosas",
            "pagina": "2",
        })

        self.assertIn(
            "socio=Las+Rosas",
            respuesta.context["query_filtros"],
        )

    def test_comparacion_metricas_historicas(self):
        """CDE-95: comparar métricas históricas entre dos años."""

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )

        self.client.force_login(self.coordinador)

        respuesta = self.client.get(
            reverse("historicos:consulta_historicos"),
            {
                "anio_a": "2023",
                "anio_b": "2024",
            },
        )

        self.assertEqual(respuesta.status_code, 200)

        comparacion = respuesta.context["comparacion"]
        self.assertIsNotNone(comparacion)

        self.assertEqual(comparacion["a"]["participaciones"], 1)
        self.assertEqual(comparacion["b"]["participaciones"], 1)

        self.assertEqual(comparacion["a"]["socios"], 1)
        self.assertEqual(comparacion["b"]["socios"], 1)

        self.assertEqual(
            comparacion["diferencia_participaciones"], 0
        )
        self.assertEqual(
            comparacion["diferencia_socios"], 0
        )

        self.assertEqual(comparacion["a"]["beneficiarios"], 40)
        self.assertEqual(comparacion["b"]["beneficiarios"], 40)

    def test_organizaciones_unicas_vs_participaciones(self):
        """CDE-96: distinguir socios únicos de participaciones."""

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )

        self.client.force_login(self.coordinador)

        respuesta = self.client.get(
            reverse("historicos:consulta_historicos")
        )

        self.assertEqual(respuesta.status_code, 200)

        self.assertEqual(
            respuesta.context["total_participaciones"], 3
        )

        self.assertEqual(
            respuesta.context["total_socios"], 2
        )
        respuesta = self.client.get(
            reverse("historicos:consulta_historicos"),
            {"socio": "Las Rosas"},
        )

        self.assertEqual(
            respuesta.context["total_participaciones"], 2
        )
        self.assertEqual(
            respuesta.context["total_socios"], 1
        )

    def test_exportar_historicos_excel_filtrado(self):
        """CDE-97: exportar únicamente registros históricos filtrados."""

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )

        self.client.force_login(self.coordinador)

        respuesta = self.client.get(
            reverse("historicos:exportar_historicos"),
            {"anio": "2024"},
        )

        self.assertEqual(respuesta.status_code, 200)

        self.assertIn(
            "spreadsheetml.sheet",
            respuesta["Content-Type"],
        )

        libro = openpyxl.load_workbook(
            BytesIO(respuesta.content)
        )

        hoja = libro.active
        self.assertEqual(hoja.max_row, 2)

        self.assertEqual(
            hoja.cell(row=1, column=1).value,
            "Socio comunitario",
        )

        self.assertEqual(
            hoja.cell(row=2, column=1).value,
            "Fundación Las Rosas",
        )
        nombres = [
            hoja.cell(row=fila, column=1).value
            for fila in range(2, hoja.max_row + 1)
        ]

        self.assertNotIn("Junta Vecinal Norte", nombres)

    def test_exportar_historicos_sin_limite_paginacion(self):
        """CDE-97: el Excel exporta todos los resultados, no solo una página."""

        carga = services.analizar_archivo(
            archivo=excel(ENCABEZADOS_REGISTRO, self.filas()),
            tipo=services.TIPO_REGISTRO_SOCIOS,
            usuario=self.coordinador,
        )

        services.confirmar_carga(
            solicitud(self.coordinador),
            carga,
            self.campus,
        )
        participacion_original = ParticipacionSocio.objects.filter(
            origen="HISTORICO"
        ).first()

        self.assertIsNotNone(participacion_original)

        for _ in range(22):
            participacion_original.pk = None
            participacion_original._state.adding = True
            participacion_original.save()

        total_esperado = ParticipacionSocio.objects.filter(
            origen="HISTORICO"
        ).count()

        self.assertEqual(total_esperado, 25)

        self.client.force_login(self.coordinador)

        respuesta = self.client.get(
            reverse("historicos:exportar_historicos")
        )

        self.assertEqual(respuesta.status_code, 200)

        libro = openpyxl.load_workbook(
            BytesIO(respuesta.content)
        )

        hoja = libro.active
        self.assertEqual(hoja.max_row, 26)


@override_settings(STORAGES=ALMACEN_EN_MEMORIA)
class ImportarFormularioTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("cargar_formularios", stdout=StringIO())
        cls.coordinador = get_user_model().objects.create_user(username="coord_form", password="x")
        cls.coordinador.groups.add(Group.objects.create(name="Coordinador"))
        sede = Sede.objects.create(nombre="Sede Santiago", activo=True)
        cls.campus = Campus.objects.create(nombre="Campus Providencia", sede=sede, activo=True)
        cls.version = services._version_publicada("EVALUACION_AS_SOCIOS")
        _, cls.preguntas = encuestas_services.cargar_estructura(cls.version)

    def fila(self, fecha):
        """Una respuesta completa con el formato de exportación de Forms."""
        valores = []
        for pregunta in self.preguntas:
            opciones = list(pregunta.opciones.all())
            if opciones:
                valores.append(opciones[0].texto)
            elif pregunta.tipo == "ESCALA":
                valores.append(5)
            elif pregunta.tipo == "RUT":
                valores.append("11.111.111-1")
            elif pregunta.tipo == "CORREO":
                valores.append("socio@ejemplo.cl")
            elif pregunta.tipo == "NUMERO":
                valores.append(3)
            elif pregunta.tipo == "FECHA":
                valores.append("10/05/2024")
            else:
                valores.append("Texto de prueba")
        return [fecha] + valores

    def encabezados(self):
        return ["Marca temporal"] + [encuestas_services.encabezado_exportacion(p) for p in self.preguntas]

    def test_importa_respuestas_historicas_con_su_fecha_y_campus(self):
        carga = services.analizar_archivo(
            archivo=excel(self.encabezados(), [self.fila("2024-06-01 09:30:00"), self.fila("")]),
            tipo="EVALUACION_AS_SOCIOS", usuario=self.coordinador,
        )
        self.assertEqual(carga.omitidas, 1)
        self.assertEqual(carga.filas.get(estado=services.FILA_OMITIDA).errores.get().codigo, "SIN_FECHA")
        self.assertEqual(carga.mapeos.count(), len(self.preguntas) + 1)

        services.confirmar_carga(solicitud(self.coordinador), carga, self.campus)
        respuesta = RespuestaFormulario.objects.get(es_historica=True)
        self.assertEqual(respuesta.origen, "HISTORICO")
        self.assertEqual(respuesta.fecha_envio.year, 2024)
        self.assertEqual(respuesta.campus, self.campus)
        self.assertIsNone(respuesta.enlace)
        self.assertEqual(respuesta.respuestas_preguntas.count(), len(self.preguntas))
        self.assertEqual(carga.filas.get(estado__in=["ACEPTADA", "OBSERVADA"]).respuesta_resultante, respuesta)

    def test_opcion_desconocida_queda_observada(self):
        pregunta = next(p for p in self.preguntas if p.opciones.exists() and not p.opciones.filter(es_otras=True).exists())
        fila = self.fila("2024-06-01 09:30:00")
        fila[self.preguntas.index(pregunta) + 1] = "Opción que no existe"
        carga = services.analizar_archivo(
            archivo=excel(self.encabezados(), [fila]), tipo="EVALUACION_AS_SOCIOS", usuario=self.coordinador,
        )
        self.assertEqual(carga.observadas, 1)
        self.assertIn("OPCION_NO_RECONOCIDA", carga.filas.get().errores.values_list("codigo", flat=True))

    def test_archivo_sin_columnas_reconocidas_se_rechaza(self):
        from django.core.exceptions import ValidationError

        with self.assertRaises(ValidationError):
            services.analizar_archivo(
                archivo=excel(["Columna A", "Columna B"], [["1", "2"]]),
                tipo="EVALUACION_AS_SOCIOS", usuario=self.coordinador,
            )
        self.assertFalse(CargaHistorica.objects.exists())


@override_settings(STORAGES=ALMACEN_EN_MEMORIA)
class PantallaImportarTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.coordinador = User.objects.create_user(username="coord_vista", password="x")
        cls.coordinador.groups.add(Group.objects.create(name="Coordinador"))
        cls.docente = User.objects.create_user(username="docente_vista", password="x")
        cls.docente.groups.add(Group.objects.create(name="Docente"))
        sede = Sede.objects.create(nombre="Sede Santiago", activo=True)
        cls.campus = Campus.objects.create(nombre="Campus Providencia", sede=sede, activo=True)

    def test_docente_no_entra(self):
        self.client.force_login(self.docente)
        respuesta = self.client.get(reverse("historicos:cargas"))
        self.assertNotEqual(respuesta.status_code, 200)

    def test_subir_revisar_y_confirmar(self):
        self.client.force_login(self.coordinador)
        self.assertContains(self.client.get(reverse("historicos:cargas")), "Arrastra aquí el Excel")

        archivo = excel(ENCABEZADOS_REGISTRO, [[
            "2023-04-01 10:00:00", "Ana", "", "Santiago", "", "Fundación Uno", "", "", "", "", "", "", "",
            "2023", "", "", "", "",
        ]])
        respuesta = self.client.post(reverse("historicos:cargas"), {"tipo": services.TIPO_REGISTRO_SOCIOS, "archivo": archivo})
        carga = CargaHistorica.objects.get()
        self.assertRedirects(respuesta, reverse("historicos:detalle", args=[carga.id]))

        detalle = self.client.get(reverse("historicos:detalle", args=[carga.id]))
        self.assertContains(detalle, "Importar 1 fila")
        self.assertContains(detalle, "Sin comuna")

        self.client.post(reverse("historicos:confirmar", args=[carga.id]), {"campus": self.campus.id})
        self.assertTrue(SocioComunitario.objects.filter(nombre_organizacion="Fundación Uno").exists())
        self.assertContains(self.client.get(reverse("historicos:detalle", args=[carga.id])), "Importada")

    def test_cancelar_no_importa(self):
        self.client.force_login(self.coordinador)
        archivo = excel(ENCABEZADOS_REGISTRO, [["2023-04-01", "Ana", "", "", "", "Fundación Dos"] + [""] * 12])
        self.client.post(reverse("historicos:cargas"), {"tipo": services.TIPO_REGISTRO_SOCIOS, "archivo": archivo})
        carga = CargaHistorica.objects.get()
        self.client.post(reverse("historicos:cancelar", args=[carga.id]))
        carga.refresh_from_db()
        self.assertEqual(carga.estado, services.ESTADO_CANCELADA)
        self.assertFalse(SocioComunitario.objects.exists())


class ValidacionEncabezadosTests(TestCase):
    """Pruebas de validación de columnas históricas CDE-89."""

    def test_rechaza_encabezados_duplicados(self):
        archivo = excel(
            ["Nombre", "NOMBRE", "Fecha"],
            [["Ana", "Pedro", "2024-05-01"]]
        )

        with self.assertRaisesMessage(
            ValidationError,
            "encabezados duplicados"
        ):
            services.leer_tabla(archivo)

    def test_acepta_encabezados_diferentes(self):
        archivo = excel(
            ["Nombre", "Correo", "Fecha"],
            [["Ana", "ana@ejemplo.cl", "2024-05-01"]]
        )

        hoja, encabezados, filas = services.leer_tabla(archivo)

        self.assertEqual(hoja, "Respuestas")
        self.assertEqual(len(encabezados), 3)
        self.assertEqual(len(filas), 1)

    def test_mapeo_columnas_registro_socios(self):
        encabezados = [
            "Marca temporal",
            "Nombre del Socio Comunitario",
            "Correo electrónico del Socio Comunitarios",
            "Sede",
            "Carrera",
            "Año en el que trabajo con el SC",
        ]

        columnas = services._columnas_registro_socios(encabezados)

        self.assertEqual(
            columnas["fecha"],
            "Marca temporal"
        )
        self.assertEqual(
            columnas["socio"],
            "Nombre del Socio Comunitario"
        )
        self.assertEqual(
            columnas["correo"],
            "Correo electrónico del Socio Comunitarios"
        )
        self.assertEqual(
            columnas["sede"],
            "Sede"
        )
        self.assertEqual(
            columnas["carrera"],
            "Carrera"
        )
        self.assertEqual(
            columnas["anios"],
            "Año en el que trabajo con el SC"
        )
    
    def test_no_confundir_correo_socio_con_informante(self):
        encabezados = [
            "Correo electrónico del Socio Comunitario",
            "Correo electrónico de quien responde",
        ]

        columnas = services._columnas_registro_socios(encabezados)

        self.assertEqual(
            columnas.get("correo"),
            "Correo electrónico del Socio Comunitario"
        )

        self.assertEqual(
            columnas.get("informante_correo"),
            "Correo electrónico de quien responde"
        )

    def test_ignora_columnas_desconocidas(self):
        encabezados = [
            "Nombre del Socio Comunitario",
            "Sede",
            "Código interno desconocido",
        ]

        columnas = services._columnas_registro_socios(encabezados)

        self.assertEqual(
            columnas.get("socio"),
            "Nombre del Socio Comunitario"
        )

        self.assertEqual(
            columnas.get("sede"),
            "Sede"
        )

        self.assertNotIn(
            "Código interno desconocido",
            columnas.values()
        )
    def test_mapeo_sin_columna_socio(self):
        encabezados = [
            "Marca temporal",
            "Sede",
            "Carrera",
            "Correo electrónico del Socio Comunitario",
        ]

        columnas = services._columnas_registro_socios(encabezados)

        self.assertNotIn("socio", columnas)
        self.assertIn("fecha", columnas)
        self.assertIn("sede", columnas)
        self.assertIn("carrera", columnas)

    def test_rechaza_archivo_sin_columna_socio(self):
        encabezados = [
            "Marca temporal",
            "Sede",
            "Carrera",
        ]

        filas = [
            (
                2,
                {
                    "Marca temporal": "2024-05-01",
                    "Sede": "Providencia",
                    "Carrera": "Ingeniería Civil Informática",
                }
            )
        ]

        mapeo, resultados = services._analizar_registro_socios(
            encabezados,
            filas
        )

        self.assertEqual(mapeo, {})
        self.assertEqual(resultados, [])

    def test_importacion_rechaza_archivo_sin_socio(self):

        archivo = excel(
            ["Marca temporal", "Sede", "Carrera"],
            [["2024-05-01", "Santiago", "Psicología"]]
        )

        with patch.object(
            services,
            "_analizar_registro_socios",
            return_value=({}, [])
        ):
            with self.assertRaisesMessage(
                ValidationError,
                "No se reconoció ninguna columna"
            ):
                services.analizar_archivo(
                    archivo=archivo,
                    tipo=services.TIPO_REGISTRO_SOCIOS,
                    usuario=None,
                )

    def test_rechaza_hoja_inexistente(self):
        archivo = excel(
            ["Nombre", "Fecha"],
            [["Socio A", "2025-01-01"]]
        )

        with self.assertRaisesMessage(
            ValidationError,
            "no existe en el archivo Excel"
        ):
            services.leer_tabla(
                archivo,
                nombre_hoja="Hoja inexistente"
            )
    def test_obtener_hojas_excel(self):
        archivo = excel(
            ["Nombre", "Fecha"],
            [["Socio A", "2025-01-01"]]
        )

        hojas = services.obtener_hojas_excel(archivo)

        self.assertEqual(hojas, ["Respuestas"])

    def test_obtener_varias_hojas_excel(self):
        libro = openpyxl.Workbook()
        libro.active.title = "Respuestas 2024"
        libro.create_sheet("Respuestas 2025")
        libro.create_sheet("Instrucciones")

        contenido = BytesIO()
        libro.save(contenido)
        contenido.seek(0)

        archivo = SimpleUploadedFile(
            "historicos.xlsx",
            contenido.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

        hojas = services.obtener_hojas_excel(archivo)

        self.assertEqual(
            hojas,
            ["Respuestas 2024", "Respuestas 2025", "Instrucciones"],
        )