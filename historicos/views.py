from pathlib import Path
from uuid import uuid4

from django.core.paginator import Paginator
from django.db.models import Q

from academico.models import Sede, Carrera, Asignatura
from socios.models import ParticipacionSocio
from socios.constantes import ORIGEN_HISTORICO

from django.conf import settings
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.files import File
from django.core.files.storage import storages
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from academico.models import Campus
from core.ordenamiento import OPCIONES_ORDEN, orden_pedido, ordenar_por_fecha
from usuarios.decorators import coordinador_required

from django.db.models import Sum
from django.db.models.functions import Coalesce
from . import services
from .models import CargaHistorica
from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# Filas aceptadas que se muestran de ejemplo en la vista previa.
MUESTRA_ACEPTADAS = 5

# Archivo Excel pendiente de selección de hoja.
SESION_EXCEL = "historicos_excel_pendiente"
CARPETA_TEMPORAL = "historicos/temporales"


def limpiar_excel_pendiente(request):
    """Elimina el archivo temporal asociado a la sesión."""
    pendiente = request.session.pop(SESION_EXCEL, None)
    if pendiente:
        storages["private"].delete(pendiente["ruta"])


@coordinador_required
def cargas(request):
    """Importar históricos y seleccionar una hoja antes del análisis."""

    if request.method == "POST" and request.POST.get("reiniciar_carga") == "1":
        limpiar_excel_pendiente(request)
        return redirect("historicos:cargas")
    
    hojas_disponibles = []
    hoja_elegida = request.POST.get("nombre_hoja", "")
    tipo_elegido = request.POST.get("tipo", "")
    pendiente = request.session.get(SESION_EXCEL)

    if request.method == "POST":
        archivo = request.FILES.get("archivo")

        try:
            if archivo is not None:
                limite = settings.PRIVATE_STORAGE_MAX_FILE_SIZE_MB * 1024 * 1024

                if archivo.size > limite:
                    raise ValidationError(
                        "El archivo supera el tamaño máximo permitido."
                    )

                hojas_disponibles = services.obtener_hojas_excel(archivo)

                if len(hojas_disponibles) > 1:
                    if tipo_elegido not in dict(services.tipos_origen()):
                        raise ValidationError(
                            "Elige qué formulario contiene el archivo."
                        )

                    limpiar_excel_pendiente(request)

                    ruta = (
                        f"{CARPETA_TEMPORAL}/"
                        f"{uuid4().hex}{Path(archivo.name).suffix.lower()}"
                    )

                    archivo.seek(0)
                    ruta = storages["private"].save(ruta, archivo)

                    request.session[SESION_EXCEL] = {
                        "ruta": ruta,
                        "nombre": Path(archivo.name).name,
                        "tipo": tipo_elegido,
                        "hojas": hojas_disponibles,
                        "usuario_id": request.user.pk,
                    }

                    messages.info(
                        request,
                        "El archivo contiene varias hojas. "
                        "Selecciona cuál deseas revisar.",
                    )

                else:
                    limpiar_excel_pendiente(request)

                    carga = services.analizar_archivo(
                        archivo=archivo,
                        tipo=tipo_elegido,
                        usuario=request.user,
                    )

                    return redirect(
                        "historicos:detalle",
                        carga_id=carga.id,
                    )

            elif pendiente and hoja_elegida:
                if pendiente.get("usuario_id") != request.user.pk:
                    raise ValidationError(
                        "La carga pendiente no pertenece a este usuario."
                    )

                hojas_disponibles = pendiente["hojas"]
                tipo_elegido = pendiente["tipo"]

                if hoja_elegida not in hojas_disponibles:
                    raise ValidationError("Selecciona una hoja válida.")

                with storages["private"].open(pendiente["ruta"], "rb") as temporal:
                    archivo_guardado = File(
                        temporal,
                        name=pendiente["nombre"],
                    )

                    carga = services.analizar_archivo(
                        archivo=archivo_guardado,
                        tipo=tipo_elegido,
                        usuario=request.user,
                        nombre_hoja=hoja_elegida,
                    )

                limpiar_excel_pendiente(request)

                return redirect(
                    "historicos:detalle",
                    carga_id=carga.id,
                )

            else:
                messages.error(
                    request,
                    "Elige o arrastra el archivo exportado.",
                )

        except ValidationError as error:
            messages.error(request, error.messages[0])

    pendiente = request.session.get(SESION_EXCEL)

    if pendiente and pendiente.get("usuario_id") == request.user.pk:
        hojas_disponibles = pendiente["hojas"]
        tipo_elegido = pendiente["tipo"]
    else:
        pendiente = None

    orden = orden_pedido(request)

    lista = ordenar_por_fecha(
        CargaHistorica.objects.select_related(
            "archivo",
            "responsable_usuario",
        ),
        "fecha_carga",
        orden,
    )[:50]

    tipos = dict(services.tipos_origen())

    for carga in lista:
        carga.nombre_tipo = tipos.get(
            carga.tipo_origen,
            carga.tipo_origen,
        )
        carga.nombre_estado = services.ESTADOS_CARGA.get(
            carga.estado,
            carga.estado,
        )

    return render(request, "historicos/cargas.html", {
        "tipos": services.tipos_origen(),
        "tipo_elegido": tipo_elegido,
        "hojas_disponibles": hojas_disponibles,
        "hoja_elegida": hoja_elegida,
        "archivo_pendiente": bool(pendiente),
        "cargas": lista,
        "orden": orden,
        "opciones_orden": OPCIONES_ORDEN,
    })


@coordinador_required
def detalle(request, carga_id):
    """Vista previa antes de confirmar (RF-MIG-01 CA1) o resultado de una carga ya hecha."""
    carga = get_object_or_404(
        CargaHistorica.objects.select_related("archivo", "responsable_usuario"), pk=carga_id,
    )
    filas = list(carga.filas.prefetch_related("errores").order_by("numero_fila"))
    con_problemas = [fila for fila in filas if fila.estado != services.FILA_ACEPTADA]
    aceptadas = [fila for fila in filas if fila.estado == services.FILA_ACEPTADA]
    mapeos = list(carga.mapeos.order_by("id"))
    reconocidas = {mapeo.columna_origen for mapeo in mapeos}
    columnas = list(filas[0].datos_raw_json) if filas else []

    campus = Campus.objects.select_related("sede").order_by("sede__nombre", "nombre")
    # Carga inicial del piloto: Sede Santiago, Campus Providencia (RN-01).
    campus_sugerido = campus.filter(nombre__icontains="providencia").first() or campus.first()

    return render(request, "historicos/detalle.html", {
        "carga": carga,
        "nombre_tipo": services.nombre_tipo(carga.tipo_origen),
        "nombre_estado": services.ESTADOS_CARGA.get(carga.estado, carga.estado),
        "es_previa": carga.estado == services.ESTADO_PREVIA,
        "es_confirmada": carga.estado == services.ESTADO_CONFIRMADA,
        "es_registro_socios": carga.tipo_origen == services.TIPO_REGISTRO_SOCIOS,
        "mapeos": mapeos,
        "no_reconocidas": [columna for columna in columnas if columna not in reconocidas],
        "con_problemas": con_problemas,
        "muestra_aceptadas": aceptadas[:MUESTRA_ACEPTADAS],
        "importables": carga.aceptadas + carga.observadas,
        "campus": campus,
        "campus_sugerido": campus_sugerido,
        "resultado": {
            "respuestas": sum(1 for fila in filas if fila.respuesta_resultante_id),
            "socios": len({fila.socio_resultante_id for fila in filas if fila.socio_resultante_id}),
        },
    })


@coordinador_required
@require_POST
def confirmar(request, carga_id):
    carga = get_object_or_404(CargaHistorica, pk=carga_id)
    campus = Campus.objects.select_related("sede").filter(pk=request.POST.get("campus") or None).first()
    try:
        creados = services.confirmar_carga(request, carga, campus)
    except ValidationError as error:
        messages.error(request, error.messages[0])
    else:
        if "respuestas" in creados:
            resumen = f"{creados['respuestas']} respuestas importadas."
        else:
            resumen = (
                f"{creados['participaciones']} participaciones importadas "
                f"({creados['socios_nuevos']} socios nuevos)."
            )
        messages.success(request, f"Carga confirmada: {resumen}")
    return redirect("historicos:detalle", carga_id=carga.id)


@coordinador_required
@require_POST
def cancelar(request, carga_id):
    carga = get_object_or_404(CargaHistorica, pk=carga_id)
    try:
        services.cancelar_carga(carga)
    except ValidationError as error:
        messages.error(request, error.messages[0])
    else:
        messages.info(request, "Carga cancelada: no se importó ninguna fila.")
    return redirect("historicos:cargas")


@coordinador_required
def consulta_historicos(request):
    """
    CDE-94: consulta de históricos con filtros combinados.
    CDE-95: comparación de métricas entre dos años.
    """

    participaciones = ParticipacionSocio.objects.filter(
        origen=ORIGEN_HISTORICO
    ).select_related(
        "socio",
        "sede",
        "campus",
        "carrera",
        "asignatura",
    )

    anio = request.GET.get("anio", "").strip()
    sede_id = request.GET.get("sede", "").strip()
    carrera_id = request.GET.get("carrera", "").strip()
    asignatura_id = request.GET.get("asignatura", "").strip()
    socio = request.GET.get("socio", "").strip()

    if anio.isdigit() and len(anio) == 4:
        participaciones = participaciones.filter(
            fecha_inicio__year__lte=int(anio),
            fecha_fin__year__gte=int(anio),
        )

    if sede_id.isdigit():
        participaciones = participaciones.filter(
            sede_id=int(sede_id)
        )

    if carrera_id.isdigit():
        participaciones = participaciones.filter(
            carrera_id=int(carrera_id)
        )

    if asignatura_id.isdigit():
        participaciones = participaciones.filter(
            asignatura_id=int(asignatura_id)
        )

    if socio:
        participaciones = participaciones.filter(
            socio__nombre_organizacion__icontains=socio
        )

    participaciones = participaciones.order_by(
        "-fecha_inicio", "-id"
    )

    total_participaciones = participaciones.count()
    total_socios = participaciones.values(
        "socio_id"
    ).distinct().count()

    anio_a = request.GET.get("anio_a", "").strip()
    anio_b = request.GET.get("anio_b", "").strip()

    comparacion = None

    if (
        anio_a.isdigit()
        and len(anio_a) == 4
        and anio_b.isdigit()
        and len(anio_b) == 4
    ):
        base_historica = ParticipacionSocio.objects.filter(
            origen=ORIGEN_HISTORICO
        )


        if sede_id.isdigit():
            base_historica = base_historica.filter(
                sede_id=int(sede_id)
            )

        if carrera_id.isdigit():
            base_historica = base_historica.filter(
                carrera_id=int(carrera_id)
            )

        if asignatura_id.isdigit():
            base_historica = base_historica.filter(
                asignatura_id=int(asignatura_id)
            )

        if socio:
            base_historica = base_historica.filter(
                socio__nombre_organizacion__icontains=socio
            )

        def metricas_periodo(anio_comparado):
            registros = base_historica.filter(
                fecha_inicio__year__lte=anio_comparado,
                fecha_fin__year__gte=anio_comparado,
            )

            return {
                "participaciones": registros.count(),
                "socios": registros.values(
                    "socio_id"
                ).distinct().count(),
                "beneficiarios": registros.aggregate(
                    total=Sum("beneficiarios_estimados")
                )["total"],
            }

        metricas_a = metricas_periodo(int(anio_a))
        metricas_b = metricas_periodo(int(anio_b))

        comparacion = {
            "anio_a": anio_a,
            "anio_b": anio_b,
            "a": metricas_a,
            "b": metricas_b,
            "diferencia_participaciones": (
                metricas_b["participaciones"]
                - metricas_a["participaciones"]
            ),
            "diferencia_socios": (
                metricas_b["socios"]
                - metricas_a["socios"]
            ),
        }

    paginador = Paginator(participaciones, 20)
    pagina = paginador.get_page(
        request.GET.get("pagina")
    )

    parametros = request.GET.copy()
    parametros.pop("pagina", None)
    query_filtros = parametros.urlencode()

    return render(request, "historicos/consulta.html", {
        "pagina": pagina,
        "total_participaciones": total_participaciones,
        "total_socios": total_socios,
        "sedes": Sede.objects.order_by("nombre"),
        "carreras": Carrera.objects.order_by("nombre"),
        "asignaturas": Asignatura.objects.order_by("nombre"),
        "query_filtros": query_filtros,
        "comparacion": comparacion,
        "filtros": {
            "anio": anio,
            "sede": sede_id,
            "carrera": carrera_id,
            "asignatura": asignatura_id,
            "socio": socio,
            "anio_a": anio_a,
            "anio_b": anio_b,
        },
    })

@coordinador_required
def exportar_historicos(request):
    """CDE-97: exportar participaciones históricas filtradas a Excel."""

    participaciones = ParticipacionSocio.objects.filter(
        origen=ORIGEN_HISTORICO
    ).select_related(
        "socio",
        "sede",
        "campus",
        "carrera",
        "asignatura",
    )
    anio = request.GET.get("anio", "").strip()
    sede_id = request.GET.get("sede", "").strip()
    carrera_id = request.GET.get("carrera", "").strip()
    asignatura_id = request.GET.get("asignatura", "").strip()
    socio = request.GET.get("socio", "").strip()

    if anio.isdigit() and len(anio) == 4:
        participaciones = participaciones.filter(
            fecha_inicio__year__lte=int(anio),
            fecha_fin__year__gte=int(anio),
        )

    if sede_id.isdigit():
        participaciones = participaciones.filter(
            sede_id=int(sede_id)
        )

    if carrera_id.isdigit():
        participaciones = participaciones.filter(
            carrera_id=int(carrera_id)
        )

    if asignatura_id.isdigit():
        participaciones = participaciones.filter(
            asignatura_id=int(asignatura_id)
        )

    if socio:
        participaciones = participaciones.filter(
            socio__nombre_organizacion__icontains=socio
        )

    participaciones = participaciones.order_by(
        "-fecha_inicio", "-id"
    )
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Históricos A+S"

    encabezados = [
        "Socio comunitario",
        "Sede",
        "Campus",
        "Carrera",
        "Asignatura",
        "Fecha de inicio",
        "Fecha de término",
        "Beneficiarios estimados",
        "Evaluación",
        "Origen",
    ]

    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = Font(
            bold=True,
            color="FFFFFF",
        )
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="B71C1C",
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
    for participacion in participaciones:
        hoja.append([
            participacion.socio.nombre_organizacion,
            participacion.sede.nombre if participacion.sede else "",
            participacion.campus.nombre if participacion.campus else "",
            participacion.carrera.nombre if participacion.carrera else "",
            participacion.asignatura.nombre if participacion.asignatura else "",
            participacion.fecha_inicio,
            participacion.fecha_fin,
            participacion.beneficiarios_estimados,
            participacion.evaluacion,
            participacion.origen,
        ])

    for fila in hoja.iter_rows(min_row=2):
        for indice in (5, 6):
            celda = fila[indice]
            if celda.value is not None:
                celda.number_format = "DD/MM/YYYY"

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    anchos = [36, 22, 24, 28, 30, 19, 19, 24, 16, 18]

    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[
            hoja.cell(row=1, column=indice).column_letter
        ].width = ancho

    contenido = BytesIO()
    libro.save(contenido)
    contenido.seek(0)

    respuesta = HttpResponse(
        contenido.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    respuesta["Content-Disposition"] = (
        'attachment; filename="historicos_as.xlsx"'
    )

    return respuesta
