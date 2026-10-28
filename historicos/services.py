"""
Carga histórica desde la plataforma (RF-MIG-01, RF-MIG-02, HU-08).

El Coordinador arrastra la exportación de un formulario antiguo (Excel de
Google o Microsoft Forms) y el portal la procesa en dos pasos:

1. analizar_archivo: guarda el archivo, reconoce las columnas y deja cada fila
   aceptada, observada u omitida con su causa (CargaHistorica en estado PREVIA).
   No crea respuestas ni socios.
2. confirmar_carga: importa las filas aceptadas y observadas, sin crear
   secciones ni equipos (RF-MIG-02), y deja la trazabilidad fila a fila.

Tipos de origen:
- REGISTRO_SOCIOS: el Registro de Socios Comunitarios (alimentó el mapa
  territorial). Crea socios y participaciones históricas.
- El código de cualquier formulario configurable publicado (evaluaciones A+S,
  resumen de implementación, fichas): crea respuestas históricas de ese
  formulario, con las columnas reconocidas por el texto de cada pregunta.
"""

import calendar
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from academico.models import Asignatura, Carrera, Sede
from archivos.services import guardar_archivo
from auditoria.services import registrar_auditoria
from encuestas import constantes as ce
from encuestas import services as encuestas_services
from encuestas.models import FormularioPlantilla, FormularioVersion, RespuestaFormulario
from socios.constantes import ORIGEN_HISTORICO
from socios.models import ClasificacionSocio, Comuna, ContactoSocio, ParticipacionSocio, SocioComunitario

from .models import CargaError, CargaFila, CargaHistorica, CargaMapeo

TIPO_REGISTRO_SOCIOS = "REGISTRO_SOCIOS"

ESTADO_PREVIA = "PREVIA"
ESTADO_CONFIRMADA = "CONFIRMADA"
ESTADO_CANCELADA = "CANCELADA"
ESTADOS_CARGA = {
    ESTADO_PREVIA: "Por confirmar",
    ESTADO_CONFIRMADA: "Importada",
    ESTADO_CANCELADA: "Cancelada",
}

FILA_ACEPTADA = "ACEPTADA"
FILA_OBSERVADA = "OBSERVADA"
FILA_OMITIDA = "OMITIDA"

EXTENSIONES = {"xlsx"}

# Columnas que Forms agrega solo (Google: «Marca temporal»; Microsoft: «Hora de finalización»).
ENCABEZADOS_FECHA = ["marca temporal", "timestamp", "hora de finalizacion", "completion time", "fecha de envio"]

# Un socio importado ya trabajó con la Universidad: el Coordinador decide
# después si vuelve a quedar disponible (RF-SC-04 CA2).
ESTADO_SOCIO_HISTORICO = "TRABAJO_TERMINADO"

# Columnas del Registro de Socios Comunitarios, reconocidas por fragmentos del
# encabezado normalizado (sin tildes). El orden importa: el primero que calza gana.
CAMPOS_REGISTRO_SOCIOS = [
    ("fecha", ENCABEZADOS_FECHA),
    ("informante_rut", ["rut"]),
    ("cargo", ["cargo del representante", "cargo"]),
    ("representante", ["representante"]),
    ("socio", ["nombre del socio", "socio comunitario institucion", "organizacion"]),
    ("telefono", ["telefono"]),
    ("correo", ["correo electronico del socio", "correo del socio"]),
    ("informante_correo", ["correo"]),
    ("informante_nombre", ["nombre completo", "quien esta respondiendo"]),
    ("informante_rol", ["como se identifica"]),
    ("sede", ["sede"]),
    ("facultad", ["facultad"]),
    ("carrera", ["carrera"]),
    ("asignatura", ["asignatura"]),
    ("comuna", ["comuna"]),
    ("direccion", ["direccion"]),
    ("clasificacion", ["clasifique", "clasificacion"]),
    ("anios", ["ano en el que", "anos en que", "ano"]),
    ("beneficiarios", ["beneficiarios"]),
    ("servicio", ["servicio"]),
    ("semestre", ["semestre"]),
    ("evaluacion", ["evalue", "evaluacion"]),
    ("comentario", ["comentario"]),
]
ETIQUETAS_REGISTRO_SOCIOS = {
    "fecha": "Fecha de envío",
    "informante_rut": "RUT de quien informa",
    "cargo": "Cargo del representante",
    "representante": "Representante del socio",
    "socio": "Nombre del socio comunitario",
    "telefono": "Teléfono del socio",
    "correo": "Correo del socio",
    "informante_correo": "Correo de quien informa",
    "informante_nombre": "Nombre de quien informa",
    "informante_rol": "Rol de quien informa",
    "sede": "Sede",
    "facultad": "Facultad",
    "carrera": "Carrera",
    "asignatura": "Asignatura",
    "comuna": "Comuna",
    "direccion": "Dirección exacta",
    "clasificacion": "Clasificación",
    "anios": "Años de trabajo",
    "beneficiarios": "Beneficiarios",
    "servicio": "Servicio realizado",
    "semestre": "Semestre",
    "evaluacion": "Evaluación (1 a 5)",
    "comentario": "Comentario",
}

# Semestre del Registro de Socios -> (mes de inicio, mes de término) de la participación.
MESES_SEMESTRE = {"otono": (3, 7), "primavera": (8, 12)}


def _n(texto):
    return encuestas_services.normalizar_encabezado(texto)


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


# ---------------------------------------------------------------------------
# Tipos de origen
# ---------------------------------------------------------------------------

def _version_publicada(codigo):
    return (
        FormularioVersion.objects
        .filter(plantilla__codigo=codigo, estado=ce.PUBLICADA)
        .select_related("plantilla")
        .order_by("-numero_version")
        .first()
    )


def tipos_origen():
    """[(código, nombre)]: el Registro de Socios y los formularios publicados."""
    tipos = [(TIPO_REGISTRO_SOCIOS, "Registro de Socios Comunitarios")]
    plantillas = (
        FormularioPlantilla.objects
        .filter(activo=True, versiones__estado=ce.PUBLICADA)
        .distinct()
        .order_by("titulo")
    )
    tipos += [(plantilla.codigo, plantilla.titulo) for plantilla in plantillas]
    return tipos


def nombre_tipo(codigo):
    return dict(tipos_origen()).get(codigo, codigo)


def obtener_hojas_excel(archivo):
    """Devuelve los nombres de las hojas disponibles en un Excel."""
    extension = Path(archivo.name or "").suffix.lower().lstrip(".")

    if extension not in EXTENSIONES:
        raise ValidationError(
            "Sube el archivo exportado en formato Excel (.xlsx)."
        )

    try:
        libro = openpyxl.load_workbook(
            archivo,
            read_only=True,
            data_only=True,
        )
    except Exception as error:
        raise ValidationError(
            "No se pudo leer el archivo. Verifica que sea un Excel válido."
        ) from error

    try:
        return list(libro.sheetnames)
    finally:
        libro.close()
        archivo.seek(0)

# ---------------------------------------------------------------------------
# Lectura del archivo
# ---------------------------------------------------------------------------


def leer_tabla(archivo, nombre_hoja=None):
    """
    (hoja, encabezados, filas) de la primera hoja con datos. Cada fila es
    (número de fila en Excel, {encabezado: celda}); se saltan las vacías.
    Si se indica nombre_hoja, solo se analiza esa hoja.
    """
    extension = Path(archivo.name or "").suffix.lower().lstrip(".")
    if extension not in EXTENSIONES:
        raise ValidationError("Sube el archivo exportado en formato Excel (.xlsx).")

    try:
        libro = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    except Exception as error:
        raise ValidationError(
            "No se pudo leer el archivo. Verifica que sea un Excel válido."
        ) from error

    try:
        if nombre_hoja is not None:
            if nombre_hoja not in libro.sheetnames:
                raise ValidationError(
                    f"La hoja '{nombre_hoja}' no existe en el archivo Excel."
                )
            hojas = [libro[nombre_hoja]]
        else:
            hojas = libro.worksheets

        for hoja in hojas:
            filas = hoja.iter_rows(values_only=True)
            encabezados = None
            datos = []

            for numero, fila in enumerate(filas, start=1):
                if encabezados is None:
                    if sum(1 for v in fila if _texto(v)) >= 2:
                        encabezados = [_texto(v) for v in fila]
                        normalizados = [
                            _n(nombre)
                            for nombre in encabezados
                            if nombre
                        ]
                        if len(normalizados) != len(set(normalizados)):
                            raise ValidationError(
                                f"La hoja '{hoja.title}' contiene "
                                "encabezados duplicados. "
                                "Corrige los nombres antes de importar."
                            )
                    continue

                if not any(_texto(v) for v in fila):
                    continue

                datos.append((
                    numero,
                    {
                        encabezado: valor
                        for encabezado, valor in zip(encabezados, fila)
                        if encabezado
                    }
                ))

            if encabezados and datos:
                return hoja.title, [e for e in encabezados if e], datos

    finally:
        libro.close()
        archivo.seek(0)

    raise ValidationError(
        "El archivo no tiene filas con datos bajo una fila de encabezados."
    )


def _raw(fila):
    """Celdas tal como venían, en texto, para conservarlas en la carga (RF-MIG-01 CA2)."""
    datos = {}
    for encabezado, valor in fila.items():
        if isinstance(valor, (datetime, date)):
            datos[encabezado] = valor.isoformat()
        else:
            datos[encabezado] = _texto(valor)
    return datos


def _fecha_hora(valor):
    if isinstance(valor, datetime):
        fecha = valor
    elif isinstance(valor, date):
        fecha = datetime.combine(valor, time(12))
    else:
        texto = _texto(valor)
        fecha = None
        for formato in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
            try:
                fecha = datetime.strptime(texto[:19], formato)
                break
            except ValueError:
                continue
        if fecha is None:
            return None
    return timezone.make_aware(fecha) if timezone.is_naive(fecha) else fecha


def _columna_fecha(encabezados):
    for encabezado in encabezados:
        if _n(encabezado) in ENCABEZADOS_FECHA:
            return encabezado
    return None


# ---------------------------------------------------------------------------
# Paso 1: análisis y vista previa
# ---------------------------------------------------------------------------

def analizar_archivo(*, archivo, tipo, usuario, nombre_hoja=None):
    """Guarda el archivo y deja la carga en estado PREVIA con cada fila clasificada."""
    if tipo not in dict(tipos_origen()):
        raise ValidationError("Elige qué formulario contiene el archivo.")

    hoja, encabezados, filas = leer_tabla(
    archivo,
    nombre_hoja=nombre_hoja,
    )
    if tipo == TIPO_REGISTRO_SOCIOS:
        mapeo, resultados = _analizar_registro_socios(encabezados, filas)
    else:
        mapeo, resultados = _analizar_formulario(tipo, encabezados, filas)

    if not mapeo:
        raise ValidationError(
            "No se reconoció ninguna columna. Revisa que el archivo sea la exportación de ese formulario."
        )

    with transaction.atomic():
        carga = CargaHistorica.objects.create(
            archivo=guardar_archivo(archivo, autor=usuario),
            responsable_usuario=usuario,
            tipo_origen=tipo,
            hoja=hoja,
            fecha_carga=timezone.now(),
            estado=ESTADO_PREVIA,
            total_filas=len(resultados),
            aceptadas=sum(1 for r in resultados if r["estado"] == FILA_ACEPTADA),
            observadas=sum(1 for r in resultados if r["estado"] == FILA_OBSERVADA),
            omitidas=sum(1 for r in resultados if r["estado"] == FILA_OMITIDA),
        )
        CargaMapeo.objects.bulk_create([
            CargaMapeo(carga=carga, columna_origen=columna[:180], campo_destino=campo[:180])
            for columna, campo in mapeo.items()
        ])
        # Filas y errores en dos inserciones en lote, no dos consultas por fila.
        filas = CargaFila.objects.bulk_create([
            CargaFila(
                carga=carga,
                numero_fila=resultado["numero"],
                datos_raw_json=resultado["raw"],
                estado=resultado["estado"],
            )
            for resultado in resultados
        ])
        CargaError.objects.bulk_create([
            CargaError(fila=fila, campo=(campo or "")[:120] or None, codigo=codigo, mensaje=mensaje)
            for fila, resultado in zip(filas, resultados)
            for campo, codigo, mensaje in resultado["errores"]
        ])
    return carga


def _resultado(numero, fila, errores, omitir=False):
    if omitir:
        estado = FILA_OMITIDA
    elif errores:
        estado = FILA_OBSERVADA
    else:
        estado = FILA_ACEPTADA
    return {"numero": numero, "raw": _raw(fila), "errores": errores, "estado": estado}


# --- Formularios configurables ----------------------------------------------

def _analizar_formulario(codigo, encabezados, filas):
    version = _version_publicada(codigo)
    bloques, preguntas = encuestas_services.cargar_estructura(version)
    columnas = encuestas_services.columnas_por_pregunta(preguntas, encabezados)
    columna_fecha = _columna_fecha(encabezados)

    mapeo = {columna: f"Pregunta: {p.texto}" for p in preguntas if (columna := columnas.get(p.id))}
    if mapeo and columna_fecha:
        mapeo[columna_fecha] = "Fecha de envío"

    ya_importadas = set(
        RespuestaFormulario.objects
        .filter(version__plantilla__codigo=codigo, es_historica=True)
        .values_list("fecha_envio", "respondente_correo", "respondente_rut", "respondente_nombre")
    )
    resultados = []
    for numero, fila in filas:
        resultado = _analizar_fila_formulario(
            numero, fila, bloques, preguntas, columnas, columna_fecha, ya_importadas,
        )
        resultados.append(resultado)
    return mapeo, resultados


def _celdas(fila, columnas):
    return {pregunta_id: fila.get(columna) for pregunta_id, columna in columnas.items()}


def _clave_respuesta(fecha, preguntas, valores):
    datos = encuestas_services.datos_respondente(preguntas, valores, {p.id for p in preguntas})
    return (fecha, datos.get("correo"), datos.get("rut") or None, datos.get("nombre"))


def _analizar_fila_formulario(numero, fila, bloques, preguntas, columnas, columna_fecha, ya_importadas):
    errores = []
    fecha = _fecha_hora(fila.get(columna_fecha)) if columna_fecha else None
    if fecha is None:
        errores.append(("Fecha de envío", "SIN_FECHA", "La fila no tiene una fecha de envío válida."))
        return _resultado(numero, fila, errores, omitir=True)

    valores, _, seleccion, avisos = encuestas_services.valores_desde_celdas(preguntas, _celdas(fila, columnas))
    if not valores:
        errores.append((None, "SIN_RESPUESTAS", "La fila no tiene respuestas reconocidas."))
        return _resultado(numero, fila, errores, omitir=True)

    clave = _clave_respuesta(fecha, preguntas, valores)
    if clave in ya_importadas:
        errores.append((None, "DUPLICADA", "Esta respuesta ya fue importada antes (misma fecha y respondente)."))
        return _resultado(numero, fila, errores, omitir=True)
    ya_importadas.add(clave)

    errores += [(pregunta.texto, codigo, mensaje) for pregunta, codigo, mensaje in avisos]
    _, visibles, _ = encuestas_services.calcular_visibilidad(bloques, preguntas, seleccion)
    for pregunta in preguntas:
        if pregunta.obligatoria and pregunta.id in visibles and pregunta.id not in valores and pregunta.tipo != ce.ARCHIVO:
            errores.append((pregunta.texto, "FALTA_OBLIGATORIA", "Falta la respuesta de una pregunta obligatoria."))
    return _resultado(numero, fila, errores)


# --- Registro de Socios Comunitarios ----------------------------------------

def _columnas_registro_socios(encabezados):
    """{campo: encabezado}: cada columna se asigna al primer campo que calza."""
    columnas = {}
    usados = set()
    for campo, fragmentos in CAMPOS_REGISTRO_SOCIOS:
        for encabezado in encabezados:
            if encabezado in usados:
                continue
            normalizado = _n(encabezado)
            if any(normalizado == f or normalizado.startswith(f + " ") or f" {f}" in f" {normalizado}" for f in fragmentos):
                columnas[campo] = encabezado
                usados.add(encabezado)
                break
    return columnas


class _Catalogos:
    """Búsquedas por nombre normalizado (sin tildes ni mayúsculas)."""

    def __init__(self):
        self.comunas = {_n(c.nombre): c for c in Comuna.objects.all()}
        self.carreras = {_n(c.nombre): c for c in Carrera.objects.all()}
        self.asignaturas = {_n(a.nombre): a for a in Asignatura.objects.all()}
        self.clasificaciones = {_n(c.nombre): c for c in ClasificacionSocio.objects.all()}
        self.sedes = {}
        for sede in Sede.objects.all():
            self.sedes[_n(sede.nombre)] = sede
            self.sedes[_n(sede.nombre).removeprefix("sede ")] = sede

    def buscar(self, catalogo, texto):
        return getattr(self, catalogo).get(_n(texto)) if _texto(texto) else None


def _anios(fila, columnas, fecha):
    texto = _texto(fila.get(columnas.get("anios")))
    anios = sorted({int(p) for p in texto.replace(",", ";").split(";") if p.strip().isdigit()})
    if not anios and fecha:
        anios = [fecha.year]
    return anios


def _entero(texto):
    texto = _texto(texto)
    try:
        return int(float(texto.replace(",", "."))) if texto else None
    except ValueError:
        return None


def _analizar_registro_socios(encabezados, filas):
    columnas = _columnas_registro_socios(encabezados)
    if "socio" not in columnas:
        return {}, []
    mapeo = {columna: ETIQUETAS_REGISTRO_SOCIOS[campo] for campo, columna in columnas.items()}
    catalogos = _Catalogos()
    existentes = {
        (socio_id, anio, carrera_id)
        for socio_id, anio, carrera_id in ParticipacionSocio.objects
        .filter(origen=ORIGEN_HISTORICO, fecha_inicio__isnull=False)
        .values_list("socio_id", "fecha_inicio__year", "carrera_id")
    }
    socios = {(_n(s.nombre_organizacion), s.comuna_id): s.id for s in SocioComunitario.objects.all()}
    vistas = set()

    resultados = []
    for numero, fila in filas:
        def valor(campo):
            return _texto(fila.get(columnas.get(campo))) if campo in columnas else ""

        errores = []
        nombre = valor("socio")
        if not nombre:
            errores.append((ETIQUETAS_REGISTRO_SOCIOS["socio"], "SIN_SOCIO", "Falta el nombre del socio comunitario."))
            resultados.append(_resultado(numero, fila, errores, omitir=True))
            continue

        fecha = _fecha_hora(fila.get(columnas["fecha"])) if "fecha" in columnas else None
        anios = _anios(fila, columnas, fecha)
        if not anios:
            errores.append((ETIQUETAS_REGISTRO_SOCIOS["anios"], "SIN_ANIO", "No se indica el año de trabajo ni la fecha de envío."))
            resultados.append(_resultado(numero, fila, errores, omitir=True))
            continue

        comuna = catalogos.buscar("comunas", valor("comuna"))
        carrera = catalogos.buscar("carreras", valor("carrera"))
        socio_id = socios.get((_n(nombre), comuna.id if comuna else None))
        nuevos = [
            anio for anio in anios
            if (socio_id, anio, carrera.id if carrera else None) not in existentes
            and (_n(nombre), comuna, anio, carrera) not in vistas
        ]
        if not nuevos:
            errores.append((None, "DUPLICADA", "Las participaciones de esta fila ya fueron importadas."))
            resultados.append(_resultado(numero, fila, errores, omitir=True))
            continue
        vistas.update((_n(nombre), comuna, anio, carrera) for anio in nuevos)

        if valor("comuna") and comuna is None:
            errores.append(("Comuna", "COMUNA_NO_RECONOCIDA", f"La comuna «{valor('comuna')}» no está registrada; el socio no aparecerá en el mapa."))
        elif comuna is None:
            errores.append(("Comuna", "SIN_COMUNA", "Sin comuna: el socio no aparecerá en el mapa."))
        elif comuna.lat_centro is None or comuna.lon_centro is None:
            errores.append(("Comuna", "COMUNA_SIN_UBICACION", f"La comuna {comuna.nombre} no tiene ubicación; no aparecerá en el mapa."))
        if valor("carrera") and carrera is None:
            errores.append(("Carrera", "CARRERA_NO_RECONOCIDA", f"La carrera «{valor('carrera')}» no está registrada."))
        if valor("sede") and catalogos.buscar("sedes", valor("sede")) is None:
            errores.append(("Sede", "SEDE_NO_RECONOCIDA", f"La sede «{valor('sede')}» no está registrada; se usará la elegida al confirmar."))
        if valor("clasificacion") and catalogos.buscar("clasificaciones", valor("clasificacion")) is None:
            errores.append(("Clasificación", "CLASIFICACION_NO_RECONOCIDA", f"«{valor('clasificacion')}» se guarda como otra clasificación."))
        evaluacion = _entero(valor("evaluacion"))
        if valor("evaluacion") and (evaluacion is None or not 1 <= evaluacion <= 5):
            errores.append(("Evaluación", "FUERA_DE_ESCALA", f"La evaluación «{valor('evaluacion')}» no está entre 1 y 5; no se guarda."))
        if valor("beneficiarios") and _entero(valor("beneficiarios")) is None:
            errores.append(("Beneficiarios", "NUMERO_INVALIDO", f"«{valor('beneficiarios')}» no es un número; no se guarda."))
        if len(nuevos) < len(anios):
            repetidos = ", ".join(str(a) for a in anios if a not in nuevos)
            errores.append(("Años de trabajo", "DUPLICADA", f"Ya estaban importados los años {repetidos}; solo se agregan los demás."))

        resultados.append(_resultado(numero, fila, errores))
    return mapeo, resultados


# ---------------------------------------------------------------------------
# Paso 2: confirmar o cancelar
# ---------------------------------------------------------------------------

def cancelar_carga(carga):
    if carga.estado != ESTADO_PREVIA:
        raise ValidationError("Esta carga ya no está pendiente de confirmación.")
    carga.estado = ESTADO_CANCELADA
    carga.save(update_fields=["estado"])


def confirmar_carga(request, carga, campus):
    """
    Importa las filas aceptadas y observadas. El campus elegido completa la
    sede y el campus de las filas que no traen la suya (RN-01).
    """
    if carga.estado != ESTADO_PREVIA:
        raise ValidationError("Esta carga ya no está pendiente de confirmación.")
    if campus is None:
        raise ValidationError("Elige el campus al que pertenecen los datos.")

    filas = list(carga.filas.exclude(estado=FILA_OMITIDA).order_by("numero_fila"))
    with transaction.atomic():
        if carga.tipo_origen == TIPO_REGISTRO_SOCIOS:
            creados = _importar_registro_socios(carga, filas, campus)
        else:
            creados = _importar_formulario(carga, filas, campus)
        carga.estado = ESTADO_CONFIRMADA
        carga.save(update_fields=["estado"])
        registrar_auditoria(
            request=request,
            entidad="CargaHistorica",
            entidad_id=carga.pk,
            accion="IMPORTACION_HISTORICA",
            valores_anteriores={"estado": ESTADO_PREVIA},
            valores_nuevos={
                "estado": ESTADO_CONFIRMADA,
                "tipo_origen": carga.tipo_origen,
                "archivo": carga.archivo.nombre_original,
                "campus": campus.nombre,
                "importadas": len(filas),
                "omitidas": carga.omitidas,
                **creados,
            },
        )
    return creados


def _importar_formulario(carga, filas, campus):
    version = _version_publicada(carga.tipo_origen)
    bloques, preguntas = encuestas_services.cargar_estructura(version)
    encabezados = list(filas[0].datos_raw_json) if filas else []
    columnas = encuestas_services.columnas_por_pregunta(preguntas, encabezados)
    columna_fecha = _columna_fecha(encabezados)
    catalogos = _Catalogos()
    pregunta_carrera = next((p for p in preguntas if _n(p.texto) == "carrera"), None)

    for fila in filas:
        datos = fila.datos_raw_json
        valores, otros, seleccion, _ = encuestas_services.valores_desde_celdas(preguntas, _celdas(datos, columnas))
        _, visibles, _ = encuestas_services.calcular_visibilidad(bloques, preguntas, seleccion)
        contexto = {"sede": campus.sede, "campus": campus}
        if pregunta_carrera and pregunta_carrera.id in columnas:
            carrera = catalogos.buscar("carreras", datos.get(columnas[pregunta_carrera.id]))
            if carrera:
                contexto["carrera"] = carrera
        respuesta = encuestas_services.guardar_respuesta(
            version, preguntas, valores, otros, visibles, seleccion,
            historica=True, fecha_envio=_fecha_hora(datos.get(columna_fecha)), contexto=contexto,
        )
        fila.respuesta_resultante = respuesta
    CargaFila.objects.bulk_update(filas, ["respuesta_resultante"])
    return {"respuestas": len(filas)}


def _fechas_participacion(anio, semestre):
    """Otoño: marzo a julio; Primavera: agosto a diciembre; Ambos o sin dato: el año académico."""
    inicio, termino = MESES_SEMESTRE.get(_n(semestre), (3, 12))
    return date(anio, inicio, 1), date(anio, termino, calendar.monthrange(anio, termino)[1])


def _importar_registro_socios(carga, filas, campus):
    columnas = _columnas_registro_socios(list(filas[0].datos_raw_json) if filas else [])
    catalogos = _Catalogos()
    otras = catalogos.buscar("clasificaciones", "Otras") or catalogos.buscar("clasificaciones", "Otra")
    socios = {(_n(s.nombre_organizacion), s.comuna_id): s for s in SocioComunitario.objects.all()}
    existentes = set(
        ParticipacionSocio.objects
        .filter(origen=ORIGEN_HISTORICO, fecha_inicio__isnull=False)
        .values_list("socio_id", "fecha_inicio__year", "carrera_id")
    )
    socios_creados = participaciones_creadas = 0

    for fila in filas:
        datos = fila.datos_raw_json

        def valor(campo):
            return _texto(datos.get(columnas.get(campo))) if campo in columnas else ""

        fecha = _fecha_hora(datos.get(columnas["fecha"])) if "fecha" in columnas else None
        comuna = catalogos.buscar("comunas", valor("comuna"))
        carrera = catalogos.buscar("carreras", valor("carrera"))
        sede = catalogos.buscar("sedes", valor("sede")) or campus.sede
        clave = (_n(valor("socio")), comuna.id if comuna else None)

        socio = socios.get(clave)
        if socio is None:
            clasificacion = catalogos.buscar("clasificaciones", valor("clasificacion"))
            primer_anio = min(_anios(datos, columnas, fecha))
            socio = SocioComunitario.objects.create(
                nombre_organizacion=valor("socio")[:220],
                clasificacion=clasificacion or (otras if valor("clasificacion") else None),
                clasificacion_otra=(valor("clasificacion")[:120] or None) if not clasificacion else None,
                comuna=comuna,
                direccion_exacta=valor("direccion")[:255] or None,
                estado_revision=ESTADO_SOCIO_HISTORICO,
                es_provisional=False,
                activo=True,
                fecha_creacion=timezone.make_aware(datetime(primer_anio, 3, 1)),
            )
            if valor("representante") or valor("correo") or valor("telefono"):
                ContactoSocio.objects.create(
                    socio=socio,
                    nombre=(valor("representante") or valor("socio"))[:180],
                    cargo=valor("cargo")[:120] or None,
                    correo=valor("correo")[:254] or None,
                    telefono=valor("telefono")[:30] or None,
                    es_principal=True,
                    activo=True,
                )
            socios[clave] = socio
            socios_creados += 1

        evaluacion = _entero(valor("evaluacion"))
        ultima = None
        for anio in _anios(datos, columnas, fecha):
            if (socio.id, anio, carrera.id if carrera else None) in existentes:
                continue
            inicio, fin = _fechas_participacion(anio, valor("semestre"))
            ultima = ParticipacionSocio.objects.create(
                socio=socio,
                sede=sede,
                campus=campus if campus.sede_id == sede.id else None,
                carrera=carrera,
                asignatura=catalogos.buscar("asignaturas", valor("asignatura")),
                servicio_realizado=valor("servicio") or None,
                beneficiarios_estimados=_entero(valor("beneficiarios")),
                evaluacion=evaluacion if evaluacion and 1 <= evaluacion <= 5 else None,
                comentario=valor("comentario") or None,
                fecha_inicio=inicio,
                fecha_fin=fin,
                origen=ORIGEN_HISTORICO,
            )
            existentes.add((socio.id, anio, carrera.id if carrera else None))
            participaciones_creadas += 1

        fila.socio_resultante = socio
        fila.participacion_resultante = ultima
    CargaFila.objects.bulk_update(filas, ["socio_resultante", "participacion_resultante"])
    return {"socios_nuevos": socios_creados, "participaciones": participaciones_creadas}
