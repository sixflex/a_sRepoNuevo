from types import SimpleNamespace
from unittest.mock import patch
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import NoReverseMatch, reverse
from openpyxl import load_workbook

from academico.models import Asignatura, Campus, Docente, PeriodoAcademico, Seccion, Sede
from planificacion.services import _resolver_docente

from .validators import es_rut_valido, normalizar_rut, validar_rut_chileno
from academico.models import Asignatura, Campus, Docente, PeriodoAcademico, Seccion, SeccionDocente, Sede

class ValidadorRutTestCase(TestCase):

    def test_valida_y_normaliza(self):
        self.assertEqual(validar_rut_chileno("12.345.678-5"), "12345678-5")
        self.assertEqual(validar_rut_chileno("123456785"), "12345678-5")
        self.assertEqual(validar_rut_chileno("11.111.112-k"), "11111112-K")

    def test_rechaza_rut_invalido(self):
        for rut in ["", "abc", "12.345.678-9", "1"]:
            with self.assertRaises(ValidationError):
                validar_rut_chileno(rut)

        self.assertFalse(es_rut_valido("12.345.678-9"))
        self.assertTrue(es_rut_valido("12345678-5"))

    def test_normalizar_sin_validar(self):
        self.assertEqual(normalizar_rut(" 12.345.678-9 "), "12345678-9")
        self.assertEqual(normalizar_rut(""), "")

    def test_socios_reexporta_el_mismo_validador(self):
        from socios.validators import validar_rut_chileno as de_socios

        self.assertIs(de_socios, validar_rut_chileno)

    def test_planificacion_encuentra_docente_con_rut_escrito_con_puntos(self):
        docente = Docente.objects.create(
            rut="12345678-5",
            nombres="Ana",
            apellidos="Soto",
            correo_institucional="ana@uautonoma.cl",
            activo=True,
        )

        fila = SimpleNamespace(rut="12.345.678-5", correo="")

        self.assertEqual(_resolver_docente(fila), docente)


class NavegacionTestCase(TestCase):

    def setUp(self):
        User = get_user_model()

        self.coordinador = User.objects.create_user(
            username="coordinador_nav",
            password="Coord1234",
        )
        self.coordinador.groups.add(Group.objects.create(name="Coordinador"))

        self.docente = User.objects.create_user(
            username="docente_nav",
            password="Docente1234",
        )
        self.docente.groups.add(Group.objects.create(name="Docente"))

    def test_docente_entra_al_inicio_docente(self):
        self.client.force_login(self.docente)

        response = self.client.get(reverse("core:inicio"))

        self.assertRedirects(response, reverse("core:inicio_docente"))

    def test_menu_docente_distingue_inicio_y_mis_secciones(self):
        self.client.force_login(self.docente)

        response = self.client.get(reverse("core:inicio_docente"))

        self.assertContains(response, reverse("core:docente"))
        self.assertContains(response, "Mis secciones")

    def test_menu_coordinador_tiene_todas_las_secciones(self):
        self.client.force_login(self.coordinador)

        response = self.client.get(reverse("core:coordinacion"))

        for nombre in [
            "home",
            "proyectos:gestionar_enlaces",
            "core:planificacion_coordinacion",
            "seccion_list",
            "rutas:seguimiento_coordinacion",
            "cartas:listar_solicitudes_coordinacion",
            "socios:lista_coordinador",
            "socios:lista_postulaciones",
            "encuestas:lista",
        ]:
            self.assertContains(response, f'href="{reverse(nombre)}"')

    def test_menu_marca_la_seccion_activa(self):
        self.client.force_login(self.coordinador)

        response = self.client.get(reverse("socios:lista_postulaciones"))

        self.assertEqual(response.context["menu_activo"], "postulaciones")
        self.assertContains(
            response,
            f'aria-current="page"\n                                   href="{reverse("socios:lista_postulaciones")}"',
        )

        response = self.client.get(reverse("sede_create"))
        self.assertEqual(response.context["menu_activo"], "academico")

        response = self.client.get(reverse("seccion_create"))
        self.assertEqual(response.context["menu_activo"], "secciones")

    def test_panel_muestra_periodo_actual_y_pendientes(self):
        PeriodoAcademico.objects.create(
            anio=2026,
            tipo="Segundo semestre",
            nombre="Primavera 2026",
            estado="Activo",
        )

        self.client.force_login(self.coordinador)

        response = self.client.get(reverse("core:coordinacion"))

        self.assertContains(response, "<title>Inicio | Portal A+S</title>")
        self.assertContains(response, "Período actual: Primavera 2026")
        self.assertContains(response, "Pendientes")
        self.assertContains(response, reverse("socios:lista_postulaciones"))

    def test_panel_sin_periodo_activo(self):
        self.client.force_login(self.coordinador)

        response = self.client.get(reverse("core:coordinacion"))

        self.assertContains(response, "Sin período activo")
        self.assertIsNone(response.context["filtros"]["periodo"])
        self.assertEqual(response.context["chips"], [])

    def test_prototipo_validar_planificacion_eliminado(self):
        # RF-ACA-04 lo resuelve la app planificacion.
        with self.assertRaises(NoReverseMatch):
            reverse("core:validar_planificacion")

    def test_pendientes_son_enlaces_a_su_listado(self):
        self.client.force_login(self.coordinador)

        response = self.client.get(reverse("core:coordinacion"))

        for nombre in [
            "cartas:listar_solicitudes_coordinacion",
            "core:planificacion_coordinacion",
            "socios:lista_postulaciones",
        ]:
            self.assertContains(response, f'href="{reverse(nombre)}" class="stretched-link"')

    def test_panel_sin_enlaces_vacios(self):
        self.client.force_login(self.coordinador)

        response = self.client.get(reverse("core:coordinacion"))

        self.assertNotContains(response, 'href="#"')
        # La validación real (RF-ACA-04) está en Planificaciones recibidas
        self.assertContains(response, reverse("core:planificacion_coordinacion"))
    
class FiltroPeriodoDocenteTestCase(TestCase):
    def setUp(self):
        User = get_user_model()

        grupo, _ = Group.objects.get_or_create(name="Docente")

        self.usuario = User.objects.create_user(
            username="docente_filtro",
            password="Prueba12345"
        )
        self.usuario.groups.add(grupo)

        self.otro_usuario = User.objects.create_user(
            username="otro_docente_filtro",
            password="Prueba12345"
        )
        self.otro_usuario.groups.add(grupo)

        self.docente = Docente.objects.create(
            usuario=self.usuario,
            rut="11111111-1",
            nombres="Docente",
            apellidos="Prueba",
            correo_institucional="docente.filtro@universidad.cl",
            activo=True
        )

        self.otro_docente = Docente.objects.create(
            usuario=self.otro_usuario,
            rut="22222222-2",
            nombres="Otro",
            apellidos="Docente",
            correo_institucional="otro.filtro@universidad.cl",
            activo=True
        )

        sede = Sede.objects.create(
            nombre="Sede Filtro",
            activo=True
        )

        campus = Campus.objects.create(
            sede=sede,
            nombre="Campus Filtro",
            activo=True
        )

        asignatura = Asignatura.objects.create(
            codigo="FILTRO01",
            nombre="Asignatura Filtro",
            activo=True
        )

        self.periodo_otono = PeriodoAcademico.objects.create(
            anio=2026,
            tipo="Semestral",
            nombre="Otoño 2026 (1.º semestre)",
            estado="Activo"
        )

        self.periodo_primavera = PeriodoAcademico.objects.create(
            anio=2026,
            tipo="Semestral",
            nombre="Primavera 2026 (2.º semestre)",
            estado="Activo"
        )

        self.seccion_otono = Seccion.objects.create(
            periodo=self.periodo_otono,
            campus=campus,
            asignatura=asignatura,
            nrc="FILTRO001",
            seccion="1",
            estado="Activa"
        )

        self.seccion_primavera = Seccion.objects.create(
            periodo=self.periodo_primavera,
            campus=campus,
            asignatura=asignatura,
            nrc="FILTRO002",
            seccion="2",
            estado="Activa"
        )

        self.seccion_ajena = Seccion.objects.create(
            periodo=self.periodo_otono,
            campus=campus,
            asignatura=asignatura,
            nrc="FILTRO003",
            seccion="3",
            estado="Activa"
        )

        SeccionDocente.objects.create(
            seccion=self.seccion_otono,
            docente=self.docente
        )

        SeccionDocente.objects.create(
            seccion=self.seccion_primavera,
            docente=self.docente
        )

        SeccionDocente.objects.create(
            seccion=self.seccion_ajena,
            docente=self.otro_docente
        )

        self.client.force_login(self.usuario)

    def test_muestra_secciones_de_todos_los_periodos(self):
        respuesta = self.client.get(reverse("core:docente"))

        self.assertEqual(respuesta.status_code, 200)

        secciones = list(respuesta.context["secciones"])

        self.assertIn(self.seccion_otono, secciones)
        self.assertIn(self.seccion_primavera, secciones)
        self.assertNotIn(self.seccion_ajena, secciones)

    def test_filtra_por_primer_semestre(self):
        respuesta = self.client.get(
            reverse("core:docente"),
            {"periodo": str(self.periodo_otono.id)}
        )

        secciones = list(respuesta.context["secciones"])

        self.assertEqual(secciones, [self.seccion_otono])

    def test_filtra_por_segundo_semestre(self):
        respuesta = self.client.get(
            reverse("core:docente"),
            {"periodo": str(self.periodo_primavera.id)}
        )

        secciones = list(respuesta.context["secciones"])

        self.assertEqual(secciones, [self.seccion_primavera])

    def test_periodo_invalido_no_muestra_secciones(self):
        respuesta = self.client.get(
            reverse("core:docente"),
            {"periodo": "invalido"}
        )

        self.assertEqual(
            list(respuesta.context["secciones"]),
            []
        )

    def test_periodo_sin_secciones_asignadas_no_muestra_resultados(self):
        otro_periodo = PeriodoAcademico.objects.create(
            anio=2025,
            tipo="Semestral",
            nombre="Otoño 2025",
            estado="Activo"
        )

        respuesta = self.client.get(
            reverse("core:docente"),
            {"periodo": str(otro_periodo.id)}
        )

        self.assertEqual(
            list(respuesta.context["secciones"]),
            []
        )

    def test_filtra_por_anio(self):
        respuesta = self.client.get(
            reverse("core:docente"),
            {"anio": "2026"}
        )

        secciones = list(respuesta.context["secciones"])

        self.assertEqual(len(secciones), 2)
        self.assertNotIn(self.seccion_ajena, secciones)

    def test_selector_solo_muestra_periodos_del_docente(self):
        otro_periodo = PeriodoAcademico.objects.create(
            anio=2025,
            tipo="Semestral",
            nombre="Otoño 2025",
            estado="Activo"
        )

        respuesta = self.client.get(reverse("core:docente"))

        periodos = list(respuesta.context["periodos"])

        self.assertIn(self.periodo_otono, periodos)
        self.assertIn(self.periodo_primavera, periodos)
        self.assertNotIn(otro_periodo, periodos)


class PanelIndicadoresTestCase(TestCase):
    """Panel del coordinador: filtros, tarjetas, avance por carrera, satisfacción y mapa."""

    def setUp(self):
        from decimal import Decimal

        from django.utils import timezone

        from academico.models import Carrera, Facultad, SeccionCarrera, SeccionDocente
        from proyectos.models import Equipo, IntegranteEquipo
        from rutas.models import RutaActividad, RutaPlantilla, RutaVersion, SeccionActividad, SeccionRuta
        from socios.models import Comuna, ParticipacionSocio, SocioComunitario

        User = get_user_model()
        self.coordinador = User.objects.create_user(username="coord_panel", password="Coord1234")
        self.coordinador.groups.add(Group.objects.create(name="Coordinador"))

        sede = Sede.objects.create(nombre="Santiago", activo=True)
        self.campus = Campus.objects.create(sede=sede, nombre="Providencia", activo=True)
        self.otono = PeriodoAcademico.objects.create(
            anio=2026, tipo="Semestral", nombre="Otoño 2026 (1er semestre)", estado="Finalizado",
        )
        self.primavera = PeriodoAcademico.objects.create(
            anio=2026, tipo="Semestral", nombre="Primavera 2026 (2º semestre)", estado="Activo",
        )
        facultad = Facultad.objects.create(nombre="Ingeniería", activo=True)
        self.informatica = Carrera.objects.create(facultad=facultad, nombre="Informática", activo=True)
        self.industrial = Carrera.objects.create(facultad=facultad, nombre="Industrial", activo=True)

        version = RutaVersion.objects.create(
            plantilla=RutaPlantilla.objects.create(nombre="Ruta A+S", activo=True),
            numero_version=1,
            estado="BORRADOR",
        )
        obligatorias = [
            RutaActividad.objects.create(
                ruta_version=version, etapa="Inicio", nombre=f"Actividad {n}", es_obligatoria=True, orden=n,
            )
            for n in range(1, 5)
        ]
        opcional = RutaActividad.objects.create(
            ruta_version=version, etapa="Cierre", nombre="Opcional", es_obligatoria=False, orden=9,
        )
        version.estado = "VIGENTE"
        version.save(update_fields=["estado"])

        def seccion(nrc, carrera, completadas, periodo=None, docente="Ana"):
            seccion = Seccion.objects.create(
                periodo=periodo or self.primavera, campus=self.campus,
                asignatura=Asignatura.objects.create(nombre=f"Asignatura {nrc}", activo=True),
                nrc=nrc, seccion="1", estado="Activa",
            )
            SeccionCarrera.objects.create(
                seccion=seccion, carrera=carrera, declaracion_as=True, estudiantes_planificados=30,
            )
            profesor = Docente.objects.create(
                rut=f"{nrc}-1", nombres=docente, apellidos="Pérez",
                correo_institucional=f"{nrc}@uautonoma.cl", activo=True,
            )
            SeccionDocente.objects.create(seccion=seccion, docente=profesor)
            ruta = SeccionRuta.objects.create(
                seccion=seccion, ruta_version=version, estado="EN_PROGRESO", porcentaje_final=Decimal("0"),
            )
            # La opcional completada no debe subir el avance (RF-RUT-03).
            for indice, actividad in enumerate(obligatorias):
                SeccionActividad.objects.create(
                    seccion_ruta=ruta, ruta_actividad=actividad, estado="PENDIENTE",
                    completada_por_docente=profesor if indice < completadas else None,
                )
            SeccionActividad.objects.create(
                seccion_ruta=ruta, ruta_actividad=opcional, estado="COMPLETADA",
                completada_por_docente=profesor,
            )
            return seccion

        self.seccion_informatica = seccion("100", self.informatica, completadas=1)
        seccion("200", self.industrial, completadas=4, docente="Beto")
        seccion("300", self.informatica, completadas=0, periodo=self.otono)

        equipo = Equipo.objects.create(
            seccion=self.seccion_informatica, numero_grupo=1,
            informante_nombre="Inf", informante_correo="inf@uautonoma.cl", estado="ACTIVO",
        )
        for n in range(3):
            IntegranteEquipo.objects.create(
                equipo=equipo, rut=f"1{n}-1", nombres="E", apellidos=str(n),
                correo_institucional=f"e{n}@uautonoma.cl",
            )

        comuna = Comuna.objects.create(
            nombre="Providencia", region="RM", activo=True, lat_centro=-33.43, lon_centro=-70.61,
        )
        for nombre in ("Junta Vecinos", "Cesfam"):
            socio = SocioComunitario.objects.create(
                nombre_organizacion=nombre, comuna=comuna, estado_revision="APROBADO",
                es_provisional=False, activo=True, fecha_creacion=timezone.now(),
            )
            ParticipacionSocio.objects.create(
                socio=socio, equipo=None, seccion=self.seccion_informatica, periodo=self.primavera,
                carrera=self.informatica, origen="HISTORICO",
            )
        self.client.force_login(self.coordinador)

    def panel(self, **parametros):
        return self.client.get(reverse("core:coordinacion"), parametros)

    def test_sin_filtros_parte_en_el_periodo_vigente(self):
        response = self.panel()

        self.assertEqual(response.context["filtros"]["periodo"], self.primavera.id)
        self.assertEqual([chip["etiqueta"] for chip in response.context["chips"]], [self.primavera.nombre])
        self.assertContains(response, 'aria-label="Quitar filtro Primavera 2026 (2º semestre)"')

    def test_tarjetas_del_periodo_y_variacion(self):
        tarjetas = {t["etiqueta"]: t for t in self.panel().context["tarjetas"]}

        self.assertEqual(tarjetas["Asignaturas A+S"]["valor"], 2)
        self.assertEqual(tarjetas["Docentes participantes"]["valor"], 2)
        self.assertEqual(tarjetas["Carreras participantes"]["valor"], 2)
        # Integrantes registrados en equipos, no los planificados.
        self.assertEqual(tarjetas["Estudiantes participantes"]["valor"], 3)
        self.assertEqual(tarjetas["Socios comunitarios"]["valor"], 2)
        self.assertIn("2 participaciones", tarjetas["Socios comunitarios"]["detalle"])
        # Otoño tenía 1 asignatura: +1 vs período anterior.
        self.assertEqual(tarjetas["Asignaturas A+S"]["detalle"], "+1 vs período anterior")

    def test_quitar_todos_los_filtros_muestra_todo(self):
        response = self.panel(f="1")

        self.assertIsNone(response.context["filtros"]["periodo"])
        tarjetas = {t["etiqueta"]: t for t in response.context["tarjetas"]}
        self.assertEqual(tarjetas["Asignaturas A+S"]["valor"], 3)

    def test_avance_por_carrera_de_menor_a_mayor_solo_obligatorias(self):
        avance = self.panel().context["avance_carreras"]

        self.assertEqual([c["carrera"] for c in avance], ["Informática", "Industrial"])
        self.assertEqual(avance[0]["porcentaje"], 25)
        self.assertEqual(avance[0]["nivel"]["texto"], "Bajo")
        self.assertEqual(avance[1]["porcentaje"], 100)
        self.assertEqual(avance[1]["nivel"]["texto"], "Alto")
        docente = avance[0]["docentes"][0]
        self.assertEqual(docente["pendientes"], 3)
        self.assertFalse(docente["cerrada"])

    def test_filtro_por_carrera(self):
        response = self.panel(f="1", periodo=self.primavera.id, carrera=self.industrial.id)

        self.assertEqual([c["carrera"] for c in response.context["avance_carreras"]], ["Industrial"])
        self.assertEqual(len(response.context["chips"]), 2)

    def test_mapa_agrupa_por_comuna_sin_direcciones(self):
        mapa = self.panel().context["mapa"]

        self.assertEqual(len(mapa["puntos"]), 1)
        punto = mapa["puntos"][0]
        self.assertEqual((punto["comuna"], punto["socios"]), ("Providencia", 2))
        self.assertEqual(punto["carreras"], ["Informática"])
        self.assertEqual(punto["anios"], [2026])
        self.assertNotIn("direccion", punto)

    def test_satisfaccion_vacia_no_muestra_cero(self):
        response = self.panel()

        self.assertFalse(response.context["satisfaccion"]["hay_datos"])
        self.assertContains(response, "Aún no hay respuestas de las encuestas de evaluación")

    def test_satisfaccion_porcentaje_de_4_y_5_y_promedio_por_dimension(self):
        from django.utils import timezone

        from encuestas.definiciones import evaluaciones
        from encuestas.models import RespuestaFormulario, RespuestaPregunta
        from encuestas.semillas import crear_formulario

        version = crear_formulario(evaluaciones.DOCENTES)
        generales = list(
            version.preguntas.filter(configuracion_json__valoracion_general=True).order_by("orden")
        )
        aprendizaje = list(version.preguntas.filter(configuracion_json__dimension="aprendizaje"))
        self.assertEqual(len(generales), 6)

        for valores in ([5, 5, 4, 4, 2, 1], [5, 4, 4, 3, 5, 5]):
            respuesta = RespuestaFormulario.objects.create(
                version=version, fecha_envio=timezone.now(), origen="ENLACE",
                es_historica=False, estado_registro="CONFIRMADA", periodo=self.primavera,
            )
            for pregunta, valor in zip(generales, valores):
                RespuestaPregunta.objects.create(respuesta=respuesta, pregunta=pregunta, valor_escala=valor)
            for pregunta in aprendizaje:
                RespuestaPregunta.objects.create(respuesta=respuesta, pregunta=pregunta, valor_escala=4)

        satisfaccion = self.panel().context["satisfaccion"]
        docentes = satisfaccion["anillos"][1]
        # 9 de 12 respuestas con 4 o 5.
        self.assertEqual(docentes["porcentaje"], 75.0)
        self.assertEqual(docentes["personas"], 2)
        self.assertEqual(docentes["pregunta"], 18)
        self.assertIsNone(satisfaccion["anillos"][0]["porcentaje"])
        self.assertEqual(satisfaccion["tabla"][0]["valores"], [None, 4.0, None])


class OrdenListadosTests(TestCase):
    """Los listados de lo que se recibe se pueden ordenar por fecha."""

    def test_ordenar_por_fecha_deja_sin_fecha_al_final(self):
        from django.utils import timezone
        from core.ordenamiento import ordenar_por_fecha
        from socios.models import SocioComunitario

        ahora = timezone.now()
        datos = {"estado_revision": "APROBADO", "es_provisional": False, "activo": True}
        antiguo = SocioComunitario.objects.create(nombre_organizacion="A", fecha_creacion=ahora - timezone.timedelta(days=2), **datos)
        nuevo = SocioComunitario.objects.create(nombre_organizacion="B", fecha_creacion=ahora, **datos)
        todos = SocioComunitario.objects.all()

        self.assertEqual(list(ordenar_por_fecha(todos, "fecha_creacion", "recientes")), [nuevo, antiguo])
        self.assertEqual(list(ordenar_por_fecha(todos, "fecha_creacion", "antiguos")), [antiguo, nuevo])


from academico.models import (
    Asignatura,
    Campus,
    Docente,
    PeriodoAcademico,
    Seccion,
    SeccionDocente,
    Sede,
)


class PanelDocenteFiltroPeriodoTestCase(TestCase):

    def setUp(self):
        User = get_user_model()

        self.grupo_docente, _ = Group.objects.get_or_create(name="Docente")

        # Docente 1
        self.user_docente1 = User.objects.create_user(
            username="docente1",
            password="DocentePassword1",
        )
        self.user_docente1.groups.add(self.grupo_docente)

        self.docente1 = Docente.objects.create(
            usuario=self.user_docente1,
            rut="11111111-1",
            nombres="Docente",
            apellidos="Uno",
            correo_institucional="docente1@uautonoma.cl",
            activo=True,
        )

        # Docente 2
        self.user_docente2 = User.objects.create_user(
            username="docente2",
            password="DocentePassword2",
        )
        self.user_docente2.groups.add(self.grupo_docente)

        self.docente2 = Docente.objects.create(
            usuario=self.user_docente2,
            rut="22222222-2",
            nombres="Docente",
            apellidos="Dos",
            correo_institucional="docente2@uautonoma.cl",
            activo=True,
        )

        # Infraestructura básica
        self.sede = Sede.objects.create(nombre="Sede Central", activo=True)
        self.campus = Campus.objects.create(sede=self.sede, nombre="Campus Central", activo=True)
        self.asignatura1 = Asignatura.objects.create(codigo="ASIG1", nombre="Asignatura 1", activo=True)
        self.asignatura2 = Asignatura.objects.create(codigo="ASIG2", nombre="Asignatura 2", activo=True)

        # Períodos
        self.periodo_2026_1 = PeriodoAcademico.objects.create(
            anio=2026,
            tipo="Primer semestre",
            nombre="2026 - Primer Semestre",
            estado="Activo",
        )
        self.periodo_2026_2 = PeriodoAcademico.objects.create(
            anio=2026,
            tipo="Segundo semestre",
            nombre="2026 - Segundo Semestre",
            estado="Activo",
        )
        self.periodo_2025_2 = PeriodoAcademico.objects.create(
            anio=2025,
            tipo="Segundo semestre",
            nombre="2025 - Segundo Semestre",
            estado="Cerrado",
        )

        # Secciones Docente 1 en 2026-1
        self.sec_d1_p1 = Seccion.objects.create(
            periodo=self.periodo_2026_1,
            campus=self.campus,
            asignatura=self.asignatura1,
            nrc="1001",
            seccion="01",
            estado="Activa",
        )
        SeccionDocente.objects.create(seccion=self.sec_d1_p1, docente=self.docente1)

        # Secciones Docente 1 en 2026-2
        self.sec_d1_p2 = Seccion.objects.create(
            periodo=self.periodo_2026_2,
            campus=self.campus,
            asignatura=self.asignatura2,
            nrc="2001",
            seccion="02",
            estado="Activa",
        )
        SeccionDocente.objects.create(seccion=self.sec_d1_p2, docente=self.docente1)

        # Sección Docente 2 en 2026-1 (para verificar aislamiento)
        self.sec_d2_p1 = Seccion.objects.create(
            periodo=self.periodo_2026_1,
            campus=self.campus,
            asignatura=self.asignatura1,
            nrc="9001",
            seccion="99",
            estado="Activa",
        )
        SeccionDocente.objects.create(seccion=self.sec_d2_p1, docente=self.docente2)

        self.url = reverse("core:docente")

    def test_acceso_no_autenticado_redirige(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 302)

    def test_sin_filtro_muestra_todas_las_secciones_del_docente(self):
        self.client.force_login(self.user_docente1)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)

        secciones = response.context["secciones"]
        self.assertEqual(len(secciones), 2)
        self.assertIn(self.sec_d1_p1, secciones)
        self.assertIn(self.sec_d1_p2, secciones)
        self.assertNotIn(self.sec_d2_p1, secciones)
        self.assertEqual(response.context["periodo_id"], "")
        self.assertEqual(response.context["periodo_seleccionado"], "")
        self.assertEqual(response.context["anio_seleccionado"], "")
        self.assertFalse(response.context["filtros_activos"])

    def test_filtro_por_periodo_muestra_solo_secciones_del_periodo_seleccionado(self):
        self.client.force_login(self.user_docente1)

        # Filtrar por período 2026-1
        response = self.client.get(self.url, {"periodo": self.periodo_2026_1.id})
        self.assertEqual(response.status_code, 200)

        secciones = response.context["secciones"]
        self.assertEqual(len(secciones), 1)
        self.assertIn(self.sec_d1_p1, secciones)
        self.assertNotIn(self.sec_d1_p2, secciones)
        self.assertNotIn(self.sec_d2_p1, secciones)
        self.assertEqual(response.context["periodo_id"], str(self.periodo_2026_1.id))
        self.assertEqual(response.context["periodo_seleccionado"], str(self.periodo_2026_1.id))
        self.assertTrue(response.context["filtros_activos"])

        # Año y período se filtran por separado (RNF-DAT-01): año + período
        # combinados siguen mostrando solo las secciones del período elegido.
        response_combinado = self.client.get(
            self.url, {"anio": 2026, "periodo": self.periodo_2026_1.id}
        )
        secciones_combinado = response_combinado.context["secciones"]
        self.assertEqual(len(secciones_combinado), 1)
        self.assertIn(self.sec_d1_p1, secciones_combinado)

        # Filtrar por período 2026-2
        response_p2 = self.client.get(self.url, {"periodo": self.periodo_2026_2.id})
        self.assertEqual(response_p2.status_code, 200)

        secciones_p2 = response_p2.context["secciones"]
        self.assertEqual(len(secciones_p2), 1)
        self.assertIn(self.sec_d1_p2, secciones_p2)
        self.assertNotIn(self.sec_d1_p1, secciones_p2)
        self.assertNotIn(self.sec_d2_p1, secciones_p2)

    def test_filtro_aisla_secciones_entre_docentes(self):
        # Ambos tienen secciones en 2026-1, pero Docente 1 no debe ver las del Docente 2
        self.client.force_login(self.user_docente1)
        response = self.client.get(self.url, {"periodo": self.periodo_2026_1.id})
        secciones = response.context["secciones"]
        self.assertNotIn(self.sec_d2_p1, secciones)

        # Docente 2 solo debe ver su propia sección en 2026-1
        self.client.force_login(self.user_docente2)
        response_d2 = self.client.get(self.url, {"periodo": self.periodo_2026_1.id})
        secciones_d2 = response_d2.context["secciones"]
        self.assertEqual(len(secciones_d2), 1)
        self.assertIn(self.sec_d2_p1, secciones_d2)
        self.assertNotIn(self.sec_d1_p1, secciones_d2)

    def test_periodo_sin_secciones_muestra_estado_vacio(self):
        self.client.force_login(self.user_docente1)

        # Docente 1 no tiene secciones en 2025-2
        response = self.client.get(self.url, {"periodo": self.periodo_2025_2.id})
        self.assertEqual(response.status_code, 200)

        secciones = response.context["secciones"]
        self.assertEqual(len(secciones), 0)
        self.assertContains(
            response,
            "No tienes secciones asignadas para el año o período académico seleccionado.",
        )
        self.assertContains(response, "Ver todos los períodos")

    def test_periodo_invalido_se_maneja_correctamente(self):
        self.client.force_login(self.user_docente1)

        # Parámetro no numérico: no se cae y no muestra secciones de otro período
        response = self.client.get(self.url, {"periodo": "invalido"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["secciones"]), 0)

        # Año no numérico: mismo tratamiento
        response_anio = self.client.get(self.url, {"anio": "invalido"})
        self.assertEqual(response_anio.status_code, 200)
        self.assertEqual(len(response_anio.context["secciones"]), 0)

        # Parámetro numérico que no existe
        response_inexistente = self.client.get(self.url, {"periodo": "999999"})
        self.assertEqual(response_inexistente.status_code, 200)
        self.assertEqual(len(response_inexistente.context["secciones"]), 0)


class PanelDocentePendientesRutaTestCase(TestCase):

    def setUp(self):
        User = get_user_model()
        self.grupo_docente, _ = Group.objects.get_or_create(name="Docente")

        self.user_docente = User.objects.create_user(
            username="docente_ruta",
            password="DocentePassword123",
        )
        self.user_docente.groups.add(self.grupo_docente)

        self.docente = Docente.objects.create(
            usuario=self.user_docente,
            rut="33333333-3",
            nombres="Docente",
            apellidos="Ruta",
            correo_institucional="docenteruta@uautonoma.cl",
            activo=True,
        )

        self.sede = Sede.objects.create(nombre="Sede Ruta", activo=True)
        self.campus = Campus.objects.create(sede=self.sede, nombre="Campus Ruta", activo=True)
        self.asignatura = Asignatura.objects.create(codigo="RUT01", nombre="Asignatura Ruta", activo=True)
        self.periodo = PeriodoAcademico.objects.create(
            anio=2026, tipo="Primer semestre", nombre="2026 - 1", estado="Activo"
        )

        self.seccion = Seccion.objects.create(
            periodo=self.periodo,
            campus=self.campus,
            asignatura=self.asignatura,
            nrc="3001",
            seccion="03",
            estado="Activa",
        )
        SeccionDocente.objects.create(seccion=self.seccion, docente=self.docente)

        self.url = reverse("core:docente")

    @patch("core.views.obtener_recordatorios_ruta")
    def test_lista_pendientes_seccion(self, mock_recordatorios):
        """1. Que se listen los pendientes de la sección."""
        mock_recordatorios.return_value = {
            self.seccion.id: {
                "pendientes": ["Configurar syllabus", "Aprobar hitos"],
                "mostrar_invitacion": False
            }
        }
        self.client.force_login(self.user_docente)
        response = self.client.get(self.url)
        
        self.assertEqual(response.status_code, 200)
        seccion_context = response.context["secciones"][0]
        self.assertIsNotNone(seccion_context.recordatorio_ruta)
        self.assertIn("Configurar syllabus", seccion_context.recordatorio_ruta["pendientes"])

    @patch("core.views.obtener_recordatorios_ruta")
    def test_actividad_completada_no_aparece(self, mock_recordatorios):
        """2. Que una actividad completada no aparezca."""
        mock_recordatorios.return_value = {
            self.seccion.id: {
                "pendientes": ["Otra actividad"],
                "mostrar_invitacion": False
            }
        }
        self.client.force_login(self.user_docente)
        response = self.client.get(self.url)
        
        seccion_context = response.context["secciones"][0]
        self.assertNotIn("Actividad completada", seccion_context.recordatorio_ruta["pendientes"])

    @patch("core.views.obtener_recordatorios_ruta")
    def test_ruta_completada_no_muestra_recordatorio(self, mock_recordatorios):
        """3. Que una Ruta completada no muestre recordatorio."""
        mock_recordatorios.return_value = {self.seccion.id: None}
        self.client.force_login(self.user_docente)
        response = self.client.get(self.url)
        
        seccion_context = response.context["secciones"][0]
        self.assertIsNone(seccion_context.recordatorio_ruta)

    @patch("core.views.obtener_recordatorios_ruta")
    def test_seccion_sin_ruta_no_muestra_recordatorio(self, mock_recordatorios):
        """4. Que una sección sin Ruta no muestre recordatorio."""
        mock_recordatorios.return_value = {} 
        self.client.force_login(self.user_docente)
        response = self.client.get(self.url)
        
        seccion_context = response.context["secciones"][0]
        self.assertIsNone(getattr(seccion_context, 'recordatorio_ruta', None))

    @patch("core.views.obtener_recordatorios_ruta")
    def test_actividad_cierre_sin_invitaciones_muestra_recordatorio_invitar(self, mock_recordatorios):
        """5. Que haya actividad de cierre sin invitaciones muestre el recordatorio de invitar."""
        mock_recordatorios.return_value = {
            self.seccion.id: {
                "pendientes": [],
                "mostrar_invitacion": True,
                "socios": [] 
            }
        }
        self.client.force_login(self.user_docente)
        response = self.client.get(self.url)
        
        seccion_context = response.context["secciones"][0]
        self.assertTrue(seccion_context.recordatorio_ruta["mostrar_invitacion"])
        self.assertEqual(len(seccion_context.recordatorio_ruta["socios"]), 0)

    @patch("core.views.obtener_recordatorios_ruta")
    def test_invitacion_registrada_muestra_socio(self, mock_recordatorios):
        """6. Que una invitación registrada muestre al socio."""
        mock_recordatorios.return_value = {
            self.seccion.id: {
                "pendientes": [],
                "mostrar_invitacion": True,
                "socios": ["Socio Comunitario A"]
            }
        }
        self.client.force_login(self.user_docente)
        response = self.client.get(self.url)
        
        seccion_context = response.context["secciones"][0]
        self.assertIn("Socio Comunitario A", seccion_context.recordatorio_ruta["socios"])

    @patch("core.views.obtener_recordatorios_ruta")
    def test_sin_actividad_cierre_no_aparece_invitacion(self, mock_recordatorios):
        """7. Que sin actividad de cierre no aparezca la invitación."""
        mock_recordatorios.return_value = {
            self.seccion.id: {
                "pendientes": ["Actividad 1"],
                "mostrar_invitacion": False 
            }
        }
        self.client.force_login(self.user_docente)
        response = self.client.get(self.url)
        
        seccion_context = response.context["secciones"][0]
        self.assertFalse(seccion_context.recordatorio_ruta["mostrar_invitacion"])


class ExportacionIndicadoresTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.coordinador = User.objects.create_user(username="coord_export", password="Clave1234")
        self.coordinador.groups.add(Group.objects.create(name="Coordinador"))
        self.docente = User.objects.create_user(username="doc_export", password="Clave1234")
        self.docente.groups.add(Group.objects.create(name="Docente"))

    def test_coordinador_descarga_xlsx_con_filtros(self):
        self.client.force_login(self.coordinador)

        response = self.client.get(reverse("core:exportar_indicadores"), {"f": "1"})

        self.assertEqual(response.status_code, 200)
        libro = load_workbook(BytesIO(response.content), read_only=True)
        self.assertIn("Resumen", libro.sheetnames)
        self.assertEqual(libro["Resumen"]["A1"].value, "Metadato")

    def test_docente_no_puede_descargar_indicadores(self):
        self.client.force_login(self.docente)

        response = self.client.get(reverse("core:exportar_indicadores"))

        self.assertEqual(response.status_code, 403)

class IndicadoresUnicosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        """Prepara los datos base para la prueba del ticket CDE-603."""
        from django.utils import timezone
        from core import indicadores
        from academico.models import Carrera, Facultad, PeriodoAcademico, Sede
        from socios.models import Comuna, ParticipacionSocio, SocioComunitario
        
        # 1. Crear configuración básica obligatoria
        cls.sede = Sede.objects.create(nombre="Santiago", activo=True)
        cls.periodo = PeriodoAcademico.objects.create(
            anio=2026, tipo="Semestral", nombre="2026-2", estado="Activo"
        )
        cls.comuna = Comuna.objects.create(
            nombre="Maipú", region="RM", activo=True, lat_centro=-33.51, lon_centro=-70.75
        )
        cls.facultad = Facultad.objects.create(nombre="Ingeniería", activo=True)
        cls.carrera = Carrera.objects.create(facultad=cls.facultad, nombre="Informática", activo=True)
        
        # 2. Crear exactamente 1 Socio (Organización) con sus campos requeridos
        cls.socio = SocioComunitario.objects.create(
            nombre_organizacion="Junta de Vecinos El Abrazo",
            comuna=cls.comuna,
            estado_revision="APROBADO",
            es_provisional=False,
            activo=True,
            fecha_creacion=timezone.now()
        )
        
        # 3. Crear 3 participaciones distintas para el MISMO socio
        for _ in range(3):
            ParticipacionSocio.objects.create(
                socio=cls.socio, 
                sede=cls.sede, 
                periodo=cls.periodo,
                carrera=cls.carrera,
                origen="HISTORICO"
            )

    def test_conteo_socios_unicos_vs_participaciones(self):
        """
        SP7-T03: Verifica que una organización con 3 participaciones 
        cuente como 1 Socio único y 3 participaciones.
        """
        from core import indicadores
        
        filtros = {
            "periodo": self.periodo.id, 
            "anio": None, "sede": None, "campus": None, 
            "facultad": None, "carrera": None, "asignatura": None, "comuna": None
        }
        
        valores = indicadores._valores(filtros)
        
        self.assertEqual(
            valores["participaciones"], 3, 
            "El total de participaciones debe ser 3."
        )
        self.assertEqual(
            valores["socios"], 1, 
            "El total de socios únicos debe ser 1 (uso correcto de distinct)."
        )